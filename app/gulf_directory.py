# -*- coding: utf-8 -*-
"""دليل مناقصات ومشاريع الخليج - FAWAZ  (الكويت + السعودية + الإمارات)

One program for every tender / practice / project that is open, about to be tendered, under
study or just awarded. Items are tracked across updates: anything new or changed (status,
closing date, winner, value) is logged in the `changes` table. The server re-runs the update
by itself every few minutes while it runs (auto_update_minutes, changeable from the window).

Saudi Arabia: Etimad visitor API (tenders.etimad.sa) - visitors only get the first 3 pages
(72 rows) of any search, so it is queried per activity, newest-first and oldest-first; the
periodic updates fill the rest as new tenders appear. Awards: category 6 for contracting and
O&M. UAE: Dubai eSupply guest list (form-post paging). News: Google News per country.

Kuwait sources:
  - CAPT open tenders (capt.gov.kw, walked back in newspaper-date windows - the plain list
    only shows the latest 50)
  - the weekly Kuwait Al-Youm gazette reports (tenders + CAPT decisions sheets)
  - Adham's weekly General Report (Urgent Summary N/L/U, PAHW, NEWS)
  - PAHW tenders site (ASP.NET postback paging)
  - Google News RSS (all Kuwaiti papers) for projects and news
  - CAPT Winning Bids file: awards of the last 12 months = projects that just started
Reuses the awards tool (capt tool.py) for normalisation, minutes search and the PDF engine.

  python "gulf directory.py" update [--only KW,SA,AE]   جمع من كل المصادر
  python "gulf directory.py" search "<نص>" [--country SA]
  python "gulf directory.py" report "<نص>" [--pick N]
  python "gulf directory.py" serve
"""
import sys, os, re, json, sqlite3, datetime, time, math, html, threading, collections, importlib.util
import urllib.request, urllib.parse
import xml.etree.ElementTree as ET
from pathlib import Path

import fawaz_config as cfg
import updater
import capt_tool as ct
BASE = cfg.GULF_DIR
ARCHIVE = cfg.ARCHIVE

DB = BASE / "directory.sqlite"
OUT = BASE / "تقارير الدليل"
LOG = BASE / "سجل التحديث.txt"
PORT = 8766


def auto_minutes():
    """minutes between automatic updates (0 = off) - read live, the window can change it"""
    try:
        return max(0, int(cfg.S.get("auto_update_minutes", 5)))
    except (TypeError, ValueError):
        return 5


COUNTRIES = {"KW": "الكويت", "SA": "السعودية", "AE": "الإمارات"}
GAZETTE_DIRS = [Path(p) for p in cfg.S.get("gazette_dirs", [])]
GAZETTE_PDF_DIR = Path(cfg.S.get("gazette_pdf_dir", ""))
ADHAM_DIR = Path(cfg.S.get("adham_reports_dir", ""))
norm, stems, words, dmy = ct.norm, ct.stems, ct.words, ct.dmy
TODAY = datetime.date.today

# ----------------------------------------------------------------- helpers
def log(msg, progress=None):
    line = f"[{datetime.datetime.now():%Y-%m-%d %H:%M}] {msg}"
    if progress:
        progress(msg)
    print(line, flush=True) if sys.stdout else None
    try:
        with open(LOG, "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except OSError:
        pass


def iso(d):
    if isinstance(d, datetime.datetime):
        d = d.date()
    return d.isoformat() if isinstance(d, datetime.date) else ""


def any_date(v):
    """datetime / 'dd/mm/yyyy' / 'yyyy-mm-dd' / 'سبتمبر 23, 2026' / excel serial -> date"""
    if v is None or v == "":
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        return datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))
    s = str(v).strip()
    m = re.search(r"(\d{4})\s*[-/]\s*(\d{1,2})\s*[-/]\s*(\d{1,2})", s)
    if m:
        try:
            return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            pass
    m = re.search(r"(\d{1,2})\s*[-/.]\s*(\d{1,2})\s*[-/.]\s*(\d{4})", s)
    if m:
        try:
            return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            pass
    return ct.parse_ar_date(s)


def numkey(tno):
    """digits of a tender number: 'و ك م /53/2024/2025' -> '53/2024/2025'"""
    return "/".join(ct.num_parts(tno))


GENERIC_ORG = set(norm(w) for w in "وزاره وزارة الهيئه هيئه العامه عامه شركه شركة الكويت الكويتيه المؤسسه مؤسسه "
                  "لشئون لشؤون ال و الجهاز المركزي الادارة".split())


def org_tokens(org):
    return set(s for s in stems(org) if s not in GENERIC_ORG and len(s) > 1)


def org_sim(a, b):
    ta, tb = org_tokens(a), org_tokens(b)
    if not ta or not tb:
        return 0.5
    return len(ta & tb) / min(len(ta), len(tb))


CATS = [("تكييف وتبريد", r"تكييف|تبريد|تهويه|تدفيه|شلر|chiller|hvac|air.?condition|cooling"),
        ("كهرباء", r"كهربا|محول|قواطع|كيبل|كابل|ضغط منخفض|ضغط متوسط|جهد|اناره|انارة|مولد|electric|transformer|cable|lighting|generator|switchgear"),
        ("تشغيل وصيانة", r"تشغيل|صيانه|maintenance|operation|facility management|facilities management|\bfm\b|o&m"),
        ("حريق وسلامة", r"حريق|انذار|اطفاء|fire|alarm|safety"),
        ("إنشاءات ومباني", r"انشاء|انجاز|مبني|مباني|ترميم|تشييد|توسعه|انشاءات|construction|building|refurbish|renovation|civil works"),
        ("طرق وبنية تحتية", r"طرق|طريق|جسر|جسور|تقاطع|صرف|مجاري|امطار|انابيب|بنيه تحتيه|بنية تحتية|road|bridge|sewer|drainage|infrastructure|pipeline"),
        ("مياه", r"مياه|تحليه|تناضح|مضخات|water|pump|desalination"),
        ("بترول وغاز", r"نفط|بترول|آبار|ابار|غاز|مصفاه|koc|knpc|kipic|kpc|rfp|oil|gas|refinery"),
        ("نظافة", r"نظافه|تنظيف|نفايات|cleaning|waste"),
        ("أمن وحراسة", r"حراسه|امن وحراسه|الامن والحراسه|security guard|guarding"),
        ("تقنية معلومات", r"حاسب|برمجيات|انظمه معلومات|نظم معلومات|تراخيص|سيبراني|شبكه الاتصالات|اتصالات|كاميرات|software|license|licence|network|\bict\b|\bit\b|cyber|server|cctv"),
        ("استشارات وتصميم", r"استشار|تصميم|اشراف|دراسه|consult|design|supervision|study"),
        ("توريد", r"^\s*(توريد|شراء)|توريد و|supply|purchase|parts|\breq\b"),
        ("خدمات وإعاشة", r"وجبات|اعاشه|تموين|ايدي عامله|عماله|مركبات|سيارات|catering|manpower|vehicle|car rental")]


def categorize(subject, org=""):
    s = norm(subject) + " || " + norm(org)
    out = [name for name, rx in CATS if re.search(rx, s)]
    return out or ["أخرى"]


def clean_minutes_subject(s):
    """'2130 (اعيد بحث) طلب طرح المناقصة رقم و ك م / 11 / 2025 / 2026 14 / 07 / 2026 توريد ...' -> 'توريد ...'"""
    s = str(s or "")
    s = re.sub(r"^\s*\d+\s*", "", s)
    s = re.sub(r"\(?\s*[اأ]عيد بحث\s*\)?", " ", s)
    s = re.sub(r"\d{1,2}\s*/\s*\d{1,2}\s*/\s*\d{4}", " ", s)
    m = re.search(r"طلب\s+(?:طرح|نشر|ترسي[ةه]|تمديد|تعديل|[اإ]لغاء|سحب)\s+(?:[^ ]+\s+){0,9}?رقم\s*", s)
    if m:
        rest = s[m.end():]
        # drop the tender number: leading tokens made of digits, slashes, dashes, brackets, short letter codes
        rest = re.sub(r"^(?:[\s()/\-.]*(?:\d+|[ء-ي]{1,3}\b|[A-Za-z]{1,4}\b))+[\s()/\-.]*", "", rest)
        s = s[:m.start()] + " " + rest
    return re.sub(r"\s+", " ", s).strip(" -/")


def fawaz_scope(subject):
    s = norm(subject)
    s = re.sub(r"مراقبه (ال)?(صيانه|اجهزه) (ال)?\w+", " ", s)       # department names, not the work
    if re.search(r"facility management|facilities management|hvac|air.?condition|chiller|mep maintenance|"
                 r"(maintenance|operation).{0,40}(electrical|mechanical|building|facilit|plant|station|fire|lift|elevator)", s):
        return True
    if re.search(r"نظافه|تنظيف|حراسه|وجبات|اعاشه|تموين|سيارات|مركبات|حاسب|برمجيات|تراخيص|قرطاسيه|تصوير مستندات|الات تصوير", s):
        return False
    # spare parts / consumables purchases are not service contracts
    if re.search(r"قطع غيار|استهلاك|مستهلكات|طلب شراء|^\s*(توريد|شراء|طلب)\b", s) and \
            not re.search(r"اعمال|تشغيل وصيانه|تشغيل و صيانه|عقد", s):
        return False
    return bool(re.search(r"تكييف|تبريد|تهويه|شلر|chiller|hvac", s) or
                (re.search(r"صيانه|تشغيل", s) and re.search(r"كهربا|ميكانيك|مباني|شامله|حريق|انذار|محطه|مرافق|منشات|مصاعد|مضخات", s)))


# ----------------------------------------------------------------- database
SCHEMA = """
CREATE TABLE IF NOT EXISTS items(
  id INTEGER PRIMARY KEY, ukey TEXT UNIQUE, country TEXT, kind TEXT, number TEXT, numkey TEXT, org TEXT,
  subject TEXT, status TEXT, publish TEXT, closing TEXT, ptm TEXT, fees TEXT, bond TEXT, ttype TEXT,
  value TEXT, winner TEXT, award_date TEXT, cats TEXT, fawaz INT, notes TEXT, url TEXT,
  official TEXT, official_label TEXT, first_seen TEXT, last_seen TEXT, changed TEXT, updated TEXT);
CREATE TABLE IF NOT EXISTS events(
  id INTEGER PRIMARY KEY, item_id INT, date TEXT, source TEXT, etype TEXT, text TEXT,
  url TEXT, file TEXT, page INT, UNIQUE(item_id, source, etype, text));
CREATE TABLE IF NOT EXISTS changes(
  id INTEGER PRIMARY KEY, item_id INT, at TEXT, field TEXT, old TEXT, new TEXT);
CREATE TABLE IF NOT EXISTS news(
  id INTEGER PRIMARY KEY, country TEXT, date TEXT, source TEXT, title TEXT, summary TEXT, url TEXT UNIQUE,
  query TEXT, cats TEXT, fawaz INT);
CREATE TABLE IF NOT EXISTS runs(id INTEGER PRIMARY KEY, started TEXT, finished TEXT, summary TEXT);
"""
BASELINE = "2000-01-01T00:00:00"      # first_seen of items loaded in a country's first update
TRACKED = {"status": "الحالة", "closing": "موعد الإقفال", "winner": "الفائز", "value": "القيمة", "award_date": "تاريخ الترسية"}


_MIGRATED = []


def db():
    con = sqlite3.connect(DB, timeout=30)
    con.row_factory = sqlite3.Row
    con.executescript(SCHEMA)
    if not _MIGRATED:                   # news.seen_at = when we first saw it (for the unread counts)
        if "seen_at" not in [r[1] for r in con.execute("PRAGMA table_info(news)")]:
            con.execute("ALTER TABLE news ADD COLUMN seen_at TEXT DEFAULT ''")
            con.commit()
        if "docs" not in [r[1] for r in con.execute("PRAGMA table_info(items)")]:
            con.execute("ALTER TABLE items ADD COLUMN docs TEXT DEFAULT ''")
            con.commit()
        con.execute("CREATE TABLE IF NOT EXISTS capt_docs(item_id INTEGER PRIMARY KEY, docs TEXT, fetched TEXT, logged INT)")
        if "archived" not in [r[1] for r in con.execute("PRAGMA table_info(capt_docs)")]:
            con.execute("ALTER TABLE capt_docs ADD COLUMN archived TEXT DEFAULT ''")
        con.execute("CREATE TABLE IF NOT EXISTS bid_sheets(key TEXT PRIMARY KEY, meeting TEXT, number TEXT, numkey TEXT, org TEXT, "
                    "subject TEXT, closing TEXT, meeting_date TEXT, rows TEXT, fetched TEXT)")
        con.commit()
        _MIGRATED.append(1)
    return con


STATUS_RANK = {"مترسية": 6, "ملغاة": 5, "أسعار معلنة": 4, "مقفولة - تحت الدراسة": 3, "مطروحة": 2,
               "اتوافق على طرحها": 1, "إعلان مسبق": 1, "عقد جاري (قرار الجهاز)": 0, "": -1}


class Store:
    """in-memory merge of all sources, written to the db at the end"""

    def __init__(self):
        self.items, self.by_num, self.events = [], collections.defaultdict(list), []

    def find(self, number, org, subject, country="KW"):
        nk = numkey(number)
        cands = self.by_num.get((country, nk), []) if nk and len(nk) >= 3 else []
        best, bs = None, 0
        for it in cands:
            s = org_sim(org, it["org"])
            if len(ct.num_parts(number)) >= 3 or len(nk) >= 6:
                s += 0.3                                # long numbers are nearly unique
            sub = len(set(stems(subject)) & set(stems(it["subject"]))) / max(1, min(len(set(stems(subject))), len(set(stems(it["subject"])))))
            s += 0.4 * sub
            if s > bs:
                best, bs = it, s
        return best if bs >= 0.75 else None

    def add(self, source, number="", org="", subject="", status="", url="", event=None, country="KW", **f):
        number, org, subject = (str(x or "").strip() for x in (number, org, subject))
        if not subject and not number:
            return None
        it = self.find(number, org, subject, country) if number else None
        if it is None:
            it = {"kind": f.pop("kind", "مناقصة"), "number": number, "numkey": numkey(number), "org": org,
                  "subject": subject, "subj_src": source, "status": status, "url": url, "sources": set(), "ev": [],
                  "country": country}
            for k in ("publish", "closing", "ptm", "fees", "bond", "ttype", "value", "winner", "award_date", "notes"):
                it[k] = ""
            it["docs"] = []
            self.items.append(it)
            if it["numkey"]:
                self.by_num[(country, it["numkey"])].append(it)
        else:
            f.pop("kind", None)
            # the tender's own title beats the wording of a board decision about it
            SUBJ_RANK = {"capt": 5, "pahw": 4, "capt-closing": 4, "gazette": 3, "awards": 3, "adham": 2, "capt-minutes": 0}
            if subject and SUBJ_RANK.get(source, 1) > SUBJ_RANK.get(it["subj_src"], 1):
                it["subject"], it["subj_src"] = subject, source
            if not it["org"] and org:
                it["org"] = org
            if not it["url"] and url:
                it["url"] = url
            if STATUS_RANK.get(status, -1) > STATUS_RANK.get(it["status"], -1):
                it["status"] = status
        it["sources"].add(source)
        for d in f.pop("docs", None) or []:          # tender documents published by the source
            if d["url"] not in {x["url"] for x in it.setdefault("docs", [])}:
                it["docs"].append(d)
        off = f.pop("official", None)                # (rank, url, label): the most direct official page wins
        if off and off[1] and off[0] > it.get("official", (0, "", ""))[0]:
            it["official"] = off
        for k, v in f.items():
            if v in (None, ""):
                continue
            if k == "closing":          # the latest closing date wins (postponements)
                if not it["closing"] or v > it["closing"]:
                    it["closing"] = v
            elif k == "notes":
                if v not in it["notes"]:
                    it["notes"] = (it["notes"] + " | " + v).strip(" |")
            elif not it.get(k):
                it[k] = v
        if event:
            it["ev"].append(dict(event, source=source))
        return it


# ----------------------------------------------------------------- source: CAPT open tenders
def parse_capt_cards(h):
    out = []
    for blk in h.split('class="content-box green tender-info"')[1:]:
        blk = blk.split('class="content-box grey"')[0]
        pairs = re.findall(r"<ul>\s*<li>([^<]+)</li>\s*<li[^>]*>(.*?)</li>\s*</ul>", blk, re.S)
        d = {ct.strip_tags(k): ct.strip_tags(v) for k, v in pairs}
        pop = re.search(r'data-popup-url="(/ar/tenders/initial-meeting-popup/[^"]+)"', blk)
        if d.get("الرقم"):
            d["_ptm_popup"] = pop.group(1) if pop else ""
            out.append(d)
    return out, max([int(x) for x in re.findall(r"[?&]page=(\d+)", h)] or [1])


def capt_window(d0, d1):
    q = f"news_paper_date_from={d0:%Y-%m-%d}&news_paper_date_to={d1:%Y-%m-%d}&form=date"
    url = f"{ct.SITE}/ar/tenders/opening-tenders/?{q}"
    cards, last = parse_capt_cards(ct.http_get(url))
    for p in range(2, min(last, 5) + 1):
        cc, _ = parse_capt_cards(ct.http_get(f"{url}&page={p}"))
        cards += cc
        time.sleep(0.2)
    full = last >= 5 and len(cards) >= 50
    return cards, full


def collect_capt(store, progress, days_back=330):
    end, n = TODAY(), 0
    # agency name -> ministry_code (the direct link to one tender needs both)
    home = ct.http_get(f"{ct.SITE}/ar/tenders/opening-tenders/")
    seg = home[home.find('name="ministry_code"'):]
    seg = seg[:seg.find("</select>")]
    codes = {ct.strip_tags(b): a for a, b in re.findall(r'<option value="([^"]+)"[^>]*>([^<]*)</option>', seg)}
    stack = []
    d = end
    while d > end - datetime.timedelta(days=days_back):
        stack.append((d - datetime.timedelta(days=13), d))
        d -= datetime.timedelta(days=14)
    seen = set()
    while stack:
        a, b = stack.pop(0)
        cards, full = capt_window(a, b)
        if full and (b - a).days >= 1:          # window capped at 50 -> split it
            mid = a + (b - a) / 2
            stack[:0] = [(a, mid), (mid + datetime.timedelta(days=1), b)]
            continue
        for c in cards:
            k = (c.get("الرقم"), c.get("الجهة"), c.get("الموضوع"))
            if k in seen:
                continue
            seen.add(k)
            closing = iso(any_date(c.get("اخر موعد للعطاء")))
            st = "مطروحة" if closing and closing >= iso(TODAY()) else ("مقفولة - تحت الدراسة" if closing else "مطروحة")
            org_name = (c.get("الجهة") or "").strip()
            code = codes.get(org_name) or next((v for k, v in codes.items() if org_sim(k, org_name) >= 0.8), "")
            direct = f"{ct.SITE}/ar/tenders/opening-tenders/?" + urllib.parse.urlencode(
                {"ministry_code": code, "tender_no": c.get("الرقم") or ""} if code else {"tender_no": c.get("الرقم") or ""})
            store.add("capt", c.get("الرقم"), c.get("الجهة"), c.get("الموضوع"), st,
                      official=(8, direct, "الجهاز المركزي للمناقصات"),
                      url=f"{ct.SITE}/ar/tenders/opening-tenders/", publish=iso(any_date(c.get("تاريخ الطلب"))),
                      closing=closing, fees=c.get("السعر", ""), bond=c.get("التأمين", ""), ttype=c.get("النوع", ""),
                      notes=c.get("ملاحظات", ""),
                      event={"date": iso(any_date(c.get("تاريخ الطلب"))), "etype": "طرح (الجهاز المركزي)",
                             "text": f"آخر موعد للعطاء {c.get('اخر موعد للعطاء', '')} - سعر الوثائق {c.get('السعر', '')} - التأمين {c.get('التأمين', '')} - {c.get('النوع', '')}",
                             "url": f"{ct.SITE}/ar/tenders/opening-tenders/"})
            n += 1
        progress(f"الجهاز المركزي: {a:%d/%m} - {b:%d/%m/%Y} ({len(cards)})")
    return n


# ----------------------------------------------------------------- source: Kuwait Al-Youm gazette reports
def gazette_files():
    files = {}
    for d in GAZETTE_DIRS:
        if not d.is_dir():
            continue
        for f in d.glob("Kuwait Al-Youm 1*.xlsx"):
            m = re.match(r"Kuwait Al-Youm (\d+) - (\d\d-\d\d-\d{4})", f.name)
            if m and m.group(1) not in files:
                files[m.group(1)] = (f, datetime.datetime.strptime(m.group(2), "%d-%m-%Y").date())
    return files


def sheet_rows(ws, header_key):
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows[:8]) if r and any(str(x or "").strip() == header_key for x in r)), None)
    if hi is None:
        return []
    head = [str(x or "").strip() for x in rows[hi]]
    return [dict(zip(head, r)) for r in rows[hi + 1:] if r and any(x not in (None, "") for x in r)]


def collect_gazette(store, progress):
    import openpyxl
    n = 0
    meetings = ct.load_meetings()          # meeting no. -> official minutes PDF on capt.gov.kw
    for issue, (f, idate) in sorted(gazette_files().items()):
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        pdf = GAZETTE_PDF_DIR / f"{issue}.pdf"
        src = f"الكويت اليوم {issue}"
        if "1-المناقصات" in wb.sheetnames:
            for r in sheet_rows(wb["1-المناقصات"], "الجهة"):
                num, org, sub = r.get("رقم المناقصة / الممارسة"), r.get("الجهة"), r.get("الموضوع")
                if not sub or sub == "الموضوع":         # repeated header rows inside the sheet
                    continue
                closing = iso(any_date(r.get("موعد الإقفال")))
                typ = str(r.get("النوع") or "")
                ptm = str(r.get("الاجتماع التمهيدي") or "")
                # oil-sector rows keep the notice itself in the pre-bid column
                if re.search(r"تأجيل|تمديد|تعديل (موعد )?الإقفال|الاقفال", ptm) and any_date(ptm):
                    closing = closing or iso(any_date(ptm))
                st = "مطروحة" if not closing or closing >= iso(TODAY()) else "مقفولة - تحت الدراسة"
                if re.search(r"إلغاء|الغاء", ptm + " " + typ):
                    st = "ملغاة"
                elif re.search(r"إعلان مسبق|pre-announcement", ptm + " " + str(num), re.I):
                    st = "إعلان مسبق"
                elif re.search(r"فتح المظاريف|فض العطاءات|جلسة التفاوض", ptm):
                    st = "مقفولة - تحت الدراسة"
                elif re.search(r"تأجيل|تمديد|ملحق", typ + " " + ptm) and not closing:
                    st = ""
                pg = int(r.get("صفحة PDF") or 0) if str(r.get("صفحة PDF") or "").isdigit() else 0
                off = (3, f"/api/file?path={urllib.parse.quote(str(pdf))}#page={pg}", f"جريدة الكويت اليوم {issue} ص {pg}") \
                    if pdf.exists() and pg else None
                store.add("gazette", num, org, sub, st, kind="ممارسة" if "ممارس" in typ + str(num) else "مناقصة",
                          official=off,
                          closing=closing, ptm=str(r.get("الاجتماع التمهيدي") or ""), fees=str(r.get("سعر الوثائق") or ""),
                          bond=str(r.get("التأمين الأولي") or ""), ttype=typ, publish=iso(idate),
                          event={"date": iso(idate), "etype": typ or "إعلان", "file": str(pdf) if pdf.exists() else "",
                                 "page": int(r.get("صفحة PDF") or 0) if str(r.get("صفحة PDF") or "").isdigit() else 0,
                                 "text": f"{src} ص {r.get('صفحة PDF') or ''}: إقفال {r.get('موعد الإقفال') or '-'} - تمهيدي {r.get('الاجتماع التمهيدي') or '-'} - وثائق {r.get('سعر الوثائق') or '-'} - تأمين {r.get('التأمين الأولي') or '-'}",
                                 "url": str(f)})
                n += 1
        if "2-الأخبار" in wb.sheetnames:
            for r in sheet_rows(wb["2-الأخبار"], "الجهة"):
                num, org, sub = r.get("رقم المناقصة"), r.get("الجهة"), r.get("الموضوع")
                itype, dec = str(r.get("نوع البند") or ""), str(r.get("قرار الجهاز") or "")
                if not sub:
                    continue
                d = norm(dec)
                ok = "موافق" in d and "عدم الموافق" not in d and "عدم موافق" not in d[:40]
                st = ""
                if re.search(r"طرح", itype) and ok:
                    st = "اتوافق على طرحها"
                elif re.search(r"ترسي", itype) and ok:
                    st = "مترسية"
                elif re.search(r"الغاء|إلغاء", itype) and ok:
                    st = "ملغاة"
                mm = re.search(r"(\d{1,3})\s*/\s*(20\d\d)", str(r.get("رقم الاجتماع") or ""))
                murl = meetings.get(f"{mm.group(2)}/{int(mm.group(1))}", {}).get("url", "") if mm else ""
                store.add("capt-minutes", num, org, clean_minutes_subject(sub) or sub, st, kind="مناقصة",
                          official=(4, murl, f"محضر الجهاز {mm.group(1)}/{mm.group(2)}") if murl else None,
                          event={"date": iso(idate), "etype": f"قرار الجهاز - {itype}",
                                 "text": f"اجتماع {r.get('رقم الاجتماع') or ''} قرار {r.get('رقم القرار') or ''}: {dec}",
                                 "file": str(pdf) if pdf.exists() else "",
                                 "page": int(r.get("صفحة PDF") or 0) if str(r.get("صفحة PDF") or "").isdigit() else 0,
                                 "url": str(f)})
                n += 1
        wb.close()
        progress(f"جريدة الكويت اليوم {issue}")
    return n


# ----------------------------------------------------------------- source: Adham general report
def latest_adham():
    best = None
    if not ADHAM_DIR.is_dir():
        return None
    for f in ADHAM_DIR.glob("*/*/*-General Report-*.xlsx"):
        if re.search(r"BACKUP|old|PRE|short|Mashaarea|only|Merged|before", f.name, re.I):
            continue
        m = re.match(r"(\d+)-General Report", f.name)
        if m:
            k = (int(m.group(1)), "Tarek" in f.name)
            if best is None or k > best[0]:
                best = (k, f)
    return best[1] if best else None


def collect_adham(store, progress):
    import openpyxl
    f = latest_adham()
    if not f:
        return 0
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    n, src = 0, f"تقرير مشاريع {f.name.split('-')[0]}"
    if "Urgent Summary" in wb.sheetnames:
        for r in sheet_rows(wb["Urgent Summary"], "الجهة"):
            t = str(r.get("T") or "").strip()
            if t not in ("N", "L", "U"):
                continue
            st = {"N": "", "L": "أسعار معلنة", "U": "مترسية"}[t]
            closing = iso(any_date(r.get("تاريخ الإقفال")))
            if t == "N":
                st = "مطروحة" if closing and closing >= iso(TODAY()) else "مقفولة - تحت الدراسة"
            low = str(r.get("Lowest Price") or "")
            store.add("adham", r.get("رقم المناقصة"), r.get("الجهة"), r.get("الموضوع"), st, closing=closing,
                      winner=low[:200] if t == "U" else "", value=str(r.get("Price") or "") if t == "U" else "",
                      notes=f"{r.get('N Code') or ''} في تقرير مشاريع",
                      event={"date": iso(any_date(r.get("Dt"))), "etype": {"N": "مطروحة", "L": "نتيجة أسعار", "U": "ترسية"}[t],
                             "text": f"{src} {r.get('N Code') or ''}: {low} {r.get('عروض أخرى / تفاصيل الترسية') or ''} {r.get('Status') or ''}".strip(),
                             "url": str(f)})
            n += 1
    if "NEWS" in wb.sheetnames:
        con = db()
        for r in sheet_rows(wb["NEWS"], "العنوان"):
            d = iso(any_date(r.get("التاريخ")))
            url = str(r.get("الرابط") or "") or f"{f}#{r.get('العنوان')}"
            con.execute("INSERT OR IGNORE INTO news(country,date,source,title,summary,url,query,cats,fawaz) VALUES('KW',?,?,?,?,?,?,?,?)",
                        (d, str(r.get("المصدر") or ""), str(r.get("العنوان") or ""), str(r.get("الملخص") or ""), url,
                         "تقرير مشاريع", ",".join(categorize(str(r.get("العنوان") or ""))), 0))
        con.commit()
        con.close()
    wb.close()
    progress(f"تقرير مشاريع: {f.name}")
    return n


# ----------------------------------------------------------------- source: PAHW
def pahw_parse(h):
    out = []
    idx = [m.start() for m in re.finditer(r'class="tenderContainer', h)] + [len(h)]
    for a, b in zip(idx, idx[1:]):
        seg = re.sub(r"<script.*?</script>", "", h[a:b], flags=re.S)
        t = re.sub(r"(\s*\|\s*)+", " | ", html.unescape(re.sub(r"<[^>]+>", " | ", seg)))
        num = re.search(r"المناقصة رقم: \| ([^|]+)\|", t)
        sub = re.search(r"المناقصة رقم: \| [^|]+\| ([^|]+)\|", t)
        stt = re.search(r"حالة المناقصة \| ([^|]+)\|", t)
        typ = re.search(r"نوع المناقصة \| ([^|]+)\|", t)
        pub = re.search(r"تاريخ الإعلان: \| ([\d/]+)", t)
        clo = re.search(r"تاريخ الإقفال: \| ([\d/]+)", t)
        env = re.search(r"تاريخ الاجتماع: \| ([\d/]+)", t)
        awd = re.search(r"تاريخ الترسية: \| ([\d/]+)", t)
        sig = re.search(r"تاريخ التوقيع: \| ([\d/]+)", t)
        can = re.search(r'class="cancelDate">([^<]+)<', seg)
        if num and sub:
            out.append({k: (v.group(1).strip() if v else "") for k, v in
                        dict(num=num, sub=sub, st=stt, typ=typ, pub=pub, clo=clo, env=env, awd=awd, sig=sig).items()} |
                       {"can": can.group(1).strip() if can else ""})
    return out


def pahw_docs(h, number):
    """public PDFs of one PAHW tender on a list page (file names start with the tender's number)"""
    m = re.search(r"(\d{3,5})\s*-\s*20\d\d", number) or re.search(r"(\d{3,5})", number)
    if not m:
        return []
    base, out = m.group(1), []
    for href in dict.fromkeys(re.findall(r"href=['\"](/Downloads/Tenders/[^'\"]+)['\"]", h)):
        name = urllib.parse.unquote(href.rsplit("/", 1)[-1])
        if re.search(rf"(?<!\d){base}(\d{{4}})?(?!\d)", name):
            label = "إقرار المشاركة" if "اقرار" in name else ("ملحق / مرفق" if re.search(rf"{base}\d{{4}}", name) else "إعلان / مرفقات المناقصة")
            d = re.search(r"-(20\d{6})-", name)
            out.append({"label": label, "url": "https://www.pahw.gov.kw" + urllib.parse.quote(href),
                        "date": f"{d.group(1)[:4]}-{d.group(1)[4:6]}-{d.group(1)[6:]}" if d else "", "src": "الرعاية السكنية"})
    return out


def collect_pahw(store, progress, max_pages=25):
    url = "https://www.pahw.gov.kw/Tenders_arabic"
    h = ct.http_get(url)
    n, page = 0, 1
    while True:
        for r in pahw_parse(h):
            st = r["st"]
            status = ("ملغاة" if r["can"] or "ملغ" in st else "مترسية" if r["awd"] or r["sig"] or "ترسي" in st
                      else "مقفولة - تحت الدراسة" if "دراس" in st or r["env"] else "مطروحة")
            closing = iso(any_date(r["clo"]))
            if status == "مطروحة" and closing and closing < iso(TODAY()):
                status = "مقفولة - تحت الدراسة"
            store.add("pahw", r["num"], "المؤسسة العامة للرعاية السكنية", r["sub"], status, url=url, docs=pahw_docs(h, r["num"]),
                      official=(6, url, "موقع الرعاية السكنية (مناقصات)"),
                      publish=iso(any_date(r["pub"])), closing=closing, ttype=r["typ"], award_date=iso(any_date(r["awd"])),
                      event={"date": iso(any_date(r["pub"])), "etype": f"الرعاية السكنية - {st}",
                             "text": f"الإعلان {r['pub']} - الإقفال {r['clo']} - فتح المظاريف {r['env'] or '-'} - الترسية {r['awd'] or '-'} - التوقيع {r['sig'] or '-'}" + (f" - ملغاة {r['can']}" if r["can"] else ""),
                             "url": url})
            n += 1
        page += 1
        m = re.search(r"__doPostBack\(&#39;(ctl00\$MainContent\$rptPaging\$ctl%02d\$lbPaging)&#39;" % (page - 1), h)
        if not m or page > max_pages:
            break
        form = {k: v for k, v in re.findall(r'name="(__[A-Z]+)" id="__[A-Z]+" value="([^"]*)"', h)}
        form.update({"__EVENTTARGET": m.group(1), "__EVENTARGUMENT": ""})
        req = urllib.request.Request(url, data=urllib.parse.urlencode(form).encode(), headers=dict(ct.UA, **{"Content-Type": "application/x-www-form-urlencoded"}))
        with urllib.request.urlopen(req, timeout=90) as resp:
            h = resp.read().decode("utf-8", "replace")
        progress(f"الرعاية السكنية: صفحة {page}")
    return n


# ----------------------------------------------------------------- source: CAPT bid opening (prices of every bidder)
def bid_text(rows):
    def val(r):
        v = re.search(r"[\d,]+(?:\.\d+)?", r[3] or "")
        try:
            return float(v.group(0).replace(",", "")) if v else 9e18
        except ValueError:
            return 9e18
    ok = sorted([r for r in rows if "مستبعد" not in (r[1] + r[2])], key=val)
    return " · ".join(f"L{k + 1} {r[0]} {r[3]}" for k, r in enumerate(ok[:6])) + (f" · مستبعد: {len(rows) - len(ok)}" if len(rows) > len(ok) else "")


def collect_capt_closing(store, progress):
    """keep the bid-opening sheets (the site keeps only the latest meetings) and attach them to the tenders"""
    import tender_docs as td
    con = db()
    known = {r[0]: r[1] for r in con.execute("SELECT key, fetched FROM bid_sheets")}
    fresh = (datetime.datetime.now() - datetime.timedelta(hours=12)).isoformat(timespec="seconds")
    try:
        pairs = td.closing_list()
    except Exception as e:                          # noqa: BLE001
        progress(f"فض العطاءات: {e}")
        pairs = []
    new = 0
    for meeting, number in pairs:
        k = f"{meeting}|{number}"
        if k in known and known[k] >= fresh:
            continue
        try:
            sh = td.closing_sheet(meeting, number)
        except Exception as e:                      # noqa: BLE001
            progress(f"فض العطاءات {number}: {e}")
            continue
        con.execute("INSERT OR REPLACE INTO bid_sheets VALUES(?,?,?,?,?,?,?,?,?,?)",
                    (k, meeting, sh["number"], numkey(sh["number"]), sh["org"], sh["subject"], sh["closing"], sh["meeting_date"],
                     json.dumps(sh["rows"], ensure_ascii=False), datetime.datetime.now().isoformat(timespec="seconds")))
        new += 1
        time.sleep(0.3)
    con.commit()
    n = 0
    for r in con.execute("SELECT * FROM bid_sheets"):
        rows = json.loads(r["rows"] or "[]")
        if not rows:
            continue
        md = iso(any_date(r["meeting_date"]))
        store.add("capt-closing", r["number"], r["org"], r["subject"], "أسعار معلنة",
                  official=(3, td.CLOSING, f"فض العطاءات - اجتماع {r['meeting']}"),
                  event={"date": md, "etype": f"فض العطاءات - اجتماع {r['meeting']}", "url": td.CLOSING,
                         "text": f"{len(rows)} عرض: " + bid_text(rows)})
        n += 1
    con.close()
    progress(f"فض العطاءات: {n} مناقصة ({new} جديد)")
    return n


def bid_sheet_for(number, org):
    con = db()
    nk = numkey(number)
    best = None
    for r in con.execute("SELECT * FROM bid_sheets WHERE numkey=?", (nk,)):
        if org_sim(org, r["org"]) < 0.6:              # every ministry restarts at 1/2025/2026
            continue
        if not best or org_sim(org, r["org"]) > org_sim(org, best["org"]):
            best = r
    con.close()
    return {k: best[k] for k in ("meeting", "number", "org", "closing", "meeting_date")} | {"rows": json.loads(best["rows"] or "[]")} if best else None


# ----------------------------------------------------------------- source: recent awards (awards tool Excel)
def refresh_awards(store, progress):
    """new CAPT awards into the awards Excel + the minutes list, so the directory is fresh on its own
    (the awards program does the same; the Excel is only rewritten when something new came)"""
    new, _last = ct.update_excel(progress)
    try:                                # new CAPT minutes: download + index (the awards program may be closed)
        got = ct.download_new_minutes(progress)
        if got:
            ct.index_minutes(progress)
            progress(f"محاضر الجهاز الجديدة: {got}")
    except Exception as e:              # noqa: BLE001 - the minutes are a bonus
        progress(f"المحاضر: {e}")
    try:
        gazette_watch(progress)
        index_gazette(progress)
    except Exception as e:              # noqa: BLE001
        progress(f"فهرسة الكويت اليوم: {e}")
    return len(new)


# ----------------------------------------------------------------- full text of the Kuwait Al-Youm issues
def gazette_pdfs():
    """issue -> best PDF in the gazette folder ("1811.pdf", "1811 نظيف.pdf" ...)"""
    best = {}
    if not GAZETTE_PDF_DIR.is_dir():
        return best
    rank = lambda f: (0 if "نظيف" in f.stem else 1 if f.stem.strip().isdigit() else 2, -f.stat().st_size)
    for f in GAZETTE_PDF_DIR.glob("*.pdf"):
        m = re.match(r"\s*(1[5-9]\d\d)\b", f.stem)
        if m and (m.group(1) not in best or rank(f) < rank(best[m.group(1)])):
            best[m.group(1)] = f
    return best


GZ_SITE = "https://kuwaitalyawm.media.gov.kw/online/MainEditions"
GZ_STATUS = BASE / "gazette status.json"


def gazette_status():
    try:
        return json.loads(GZ_STATUS.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def gazette_watch(progress):
    """every Sunday a new Kuwait Al-Youm issue: notice it on the official site (public number + date),
    pick its PDF up as soon as it is saved (the download needs the owner's login on the site) and index it"""
    if not GAZETTE_PDF_DIR.is_dir():
        return {}
    st = gazette_status()
    now = datetime.datetime.now()
    if not st.get("checked") or (now - datetime.datetime.fromisoformat(st["checked"])).total_seconds() > 1800:
        try:
            h = ct.http_get(GZ_SITE, tries=2)
            m = re.search(r"رقم الإصدار الحالي.*?(\d{4})", re.sub(r"<[^>]+>", " ", h), re.S)
            d = re.search(r"تاريخ الإصدار.*?(\d{1,2}/\d{1,2}/\d{4})", re.sub(r"<[^>]+>", " ", h), re.S)
            if m:
                st.update(current=m.group(1), date=d.group(1) if d else "")
        except Exception as e:                      # noqa: BLE001
            progress(f"موقع الكويت اليوم: {e}")
        st["checked"] = now.isoformat(timespec="seconds")
    pick_up_gazette_downloads(progress)
    st["have"] = bool(st.get("current") and st["current"] in gazette_pdfs())
    if st.get("current") and not st["have"]:        # with the subscription: the program gets the issue itself
        import gazette_online as go
        if all(go.account()) and (not st.get("tried") or (now - datetime.datetime.fromisoformat(st["tried"])).total_seconds() > 900):
            st["tried"] = now.isoformat(timespec="seconds")
            try:
                f = go.fetch_issue_pdf(st["current"], GAZETTE_PDF_DIR)
                st["have"], st["error"] = True, ""
                progress(f"اتجاب عدد الكويت اليوم {st['current']} بالاشتراك ({f.name})")
            except Exception as e:                  # noqa: BLE001
                st["error"] = str(e)[:200]
                progress(f"الكويت اليوم بالاشتراك: {e}")
    if st.get("current") and st.get("have"):          # fill earlier issues that never got here (one per update)
        import gazette_online as go
        if all(go.account()):
            pdfs = gazette_pdfs()
            for n in range(int(st["current"]) - 1, int(st["current"]) - 9, -1):
                if str(n) not in pdfs and str(n) not in st.get("missing_failed", []):
                    try:
                        go.fetch_issue_pdf(n, GAZETTE_PDF_DIR)
                        progress(f"اتجاب عدد الكويت اليوم {n} (كان ناقص)")
                    except Exception as e:          # noqa: BLE001
                        st.setdefault("missing_failed", []).append(str(n))
                        progress(f"العدد {n}: {e}")
                    break
    GZ_STATUS.write_text(json.dumps(st, ensure_ascii=False), encoding="utf-8")
    return st


def pick_up_gazette_downloads(progress):
    """a Kuwait Al-Youm PDF saved in the gazette folder or in Downloads under any name -> <issue>.pdf"""
    import pymupdf, shutil
    have = gazette_pdfs()
    recent = time.time() - 14 * 86400
    for folder in (GAZETTE_PDF_DIR, Path.home() / "Downloads"):
        if not folder.is_dir():
            continue
        for f in folder.glob("*.pdf"):
            try:
                if f.stat().st_mtime < recent or f.stat().st_size < 1_500_000 or re.match(r"\s*1[5-9]\d\d\b", f.stem):
                    continue
                d = pymupdf.open(f)
                txt = " ".join(d[i].get_text() for i in range(min(3, d.page_count)))
                d.close()
            except Exception:                       # noqa: BLE001
                continue
            m = re.search(r"العدد\s*(1[5-9]\d\d)", txt)
            if m and "الكويت اليوم" in txt and m.group(1) not in have:
                dst = GAZETTE_PDF_DIR / f"{m.group(1)}.pdf"
                shutil.copy2(f, dst)
                have[m.group(1)] = dst
                progress(f"لقيت عدد الكويت اليوم {m.group(1)} في {f.parent.name} واتنقل لفولدر الجريدة")


def gz_db():
    con = sqlite3.connect(BASE / "gazette index.sqlite", timeout=30)
    con.row_factory = sqlite3.Row
    con.execute("CREATE TABLE IF NOT EXISTS files(issue TEXT PRIMARY KEY, path TEXT, sig TEXT, pages INT)")
    con.execute("CREATE VIRTUAL TABLE IF NOT EXISTS pages USING fts5(text, issue UNINDEXED, page UNINDEXED, path UNINDEXED, tokenize='unicode61')")
    return con


GZ_INDEX_V = "v2"          # v2: the subscriber's watermark (company name on every page) is removed


def gazette_page_texts(d):
    """page texts without what repeats on most pages - the subscriber watermark ("شركة فواز للتجارة ..." +
    the user name on every page of a downloaded issue) and the running header; otherwise every page
    would look like news about Fawaz"""
    import collections
    import gazette_online as go
    raw = [d[i].get_text() for i in range(d.page_count)]
    cnt = collections.Counter()
    for t in raw:
        cnt.update({l.strip() for l in t.splitlines() if l.strip()})
    limit = max(4, d.page_count * 0.4)
    common = {l for l, c in cnt.items() if c >= limit}
    user = (go.account()[0] or "").strip().lower()
    out = []
    for t in raw:
        keep = [l for l in t.splitlines() if l.strip() and l.strip() not in common and (not user or l.strip().lower() != user)]
        out.append("\n".join(keep))
    return out


def index_gazette(progress, keep=40):
    """index the text of the latest issues (one-off per issue, a few seconds each)"""
    import pymupdf
    pdfs = gazette_pdfs()
    con = gz_db()
    known = {r["issue"]: r["sig"] for r in con.execute("SELECT issue, sig FROM files")}
    n = 0
    for issue in sorted(pdfs, key=int)[-keep:]:
        f = pdfs[issue]
        sig = f"{GZ_INDEX_V}|{f.name}|{f.stat().st_size}|{int(f.stat().st_mtime)}"
        if known.get(issue) == sig:
            continue
        d = pymupdf.open(f)
        con.execute("DELETE FROM pages WHERE issue=?", (issue,))
        con.executemany("INSERT INTO pages(text, issue, page, path) VALUES(?,?,?,?)",
                        [(norm(t), issue, i + 1, str(f)) for i, t in enumerate(gazette_page_texts(d))])
        con.execute("INSERT OR REPLACE INTO files VALUES(?,?,?,?)", (issue, str(f), sig, d.page_count))
        con.commit()
        d.close()
        n += 1
        progress(f"اتفهرس عدد الكويت اليوم {issue}")
    con.close()
    return n


def gazette_search(q, limit=60, issue=""):
    words = [w for w in words_of(q)]
    if not words:
        return []
    con = gz_db()
    match = " AND ".join('"' + w.replace('"', "") + '"' + ("*" if len(w) >= 4 else "") for w in words)
    rows = con.execute("SELECT issue, page, path, snippet(pages, 0, '[', ']', ' … ', 18) AS snip, bm25(pages) AS sc FROM pages "
                       "WHERE pages MATCH ?" + (" AND issue=?" if issue else "") + " ORDER BY CAST(issue AS INT) DESC, sc LIMIT ?",
                       (match, issue, limit) if issue else (match, limit)).fetchall()
    con.close()
    return [{"issue": r["issue"], "page": r["page"], "file": r["path"], "snippet": r["snip"],
             "url": f"/api/file?path={urllib.parse.quote(r['path'])}#page={r['page']}"} for r in rows]


def words_of(q):
    return [w for w in re.findall(r"[0-9a-z\u0621-\u064a]+", norm(q)) if len(w) > 1]


def collect_awards(store, progress, days=365):
    since = iso(TODAY() - datetime.timedelta(days=days))
    n = 0
    for r in ct.read_excel_rows():
        if r["date"] < since:
            continue
        wb_url = f"{ct.SITE}/ar/tenders/winning-bids/?meeting_date_from={r['date']}&meeting_date_to={r['date']}&form=date"
        store.add("awards", r["tno"], r["org"], r["subject"], "مترسية", winner=r["winner"], value=r["total"],
                  official=(5, wb_url, f"ترسيات الجهاز المركزي يوم {dmy(any_date(r['date']))}"),
                  award_date=r["date"], event={"date": r["date"], "etype": "ترسية (كشف الترسيات)",
                                               "text": f"{r['winner'] or '-'} - {r['total'] or '-'} {r['notes'] or ''}".strip(),
                                               "url": str(ct.XLSX)})
        n += 1
    progress(f"الترسيات من آخر سنة: {n}")
    return n


# ----------------------------------------------------------------- source: Saudi Arabia - Etimad
ETIMAD = "https://tenders.etimad.sa"
ET_H = {"User-Agent": ct.UA["User-Agent"], "Accept": "application/json, text/plain, */*",
        "Accept-Language": "ar,en;q=0.9", "X-Requested-With": "XMLHttpRequest",
        "Referer": ETIMAD + "/Tender/AllTendersForVisitor"}


def etimad_json(path, params=None):
    url = ETIMAD + path + ("?" + urllib.parse.urlencode(params) if params else "")
    for k in range(3):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers=ET_H), timeout=60) as r:
                return json.loads(r.read())
        except Exception:               # noqa: BLE001 - retry, the API hiccups
            time.sleep(2 * (k + 1))
    return None


def et_add(store, r, status, activity):
    ref = str(r.get("referenceNumber") or r.get("tenderId") or "")
    org = (r.get("agencyName") or "").strip()
    br = (r.get("branchName") or "").strip()
    if br and br != org and br not in org:
        org = f"{org} - {br}"
    fees = r.get("condetionalBookletPrice") or r.get("buyingCost") or r.get("invitationCost") or 0
    closing = (r.get("lastOfferPresentationDate") or "")[:10]
    url = f"{ETIMAD}/Tender/DetailsForVisitor?STenderId={urllib.parse.quote(r.get('tenderIdString') or '')}"
    ev = {"date": (r.get("submitionDate") or "")[:10], "url": url,
          "etype": "إعلان الترسية (اعتماد)" if status == "مترسية" else "طرح (منصة اعتماد)",
          "text": f"{r.get('tenderTypeName') or ''} - النشاط: {activity} - آخر موعد للعروض {closing or '-'}"
                  f" - قيمة الوثائق {fees} ريال - رقم المنافسة {r.get('tenderNumber') or '-'}"}
    store.add("etimad", ref, org, r.get("tenderName"), status, url=url, country="SA", kind="منافسة",
              official=(9, url, "منصة اعتماد"),
              publish=(r.get("submitionDate") or "")[:10], closing=closing if status != "مترسية" else "",
              fees=f"{fees:,.0f} ريال" if isinstance(fees, (int, float)) and fees else "",
              ttype=r.get("tenderTypeName") or "", notes=f"النشاط: {activity}", event=ev)


def collect_etimad(store, progress):
    acts = etimad_json("/Tender/GetMainActivitiesAsync") or []
    n = 0
    for a in acts:
        aid, name = a.get("value"), (a.get("text") or "").strip()
        if not aid or not name:
            continue
        plans = [(2, "مطروحة")]
        if aid in ("2", "3"):                      # contracting + O&M: recent awards too
            plans.append((6, "مترسية"))
        for cat, status in plans:
            base = {"PageSize": 24, "IsSearch": "true", "TenderCategory": cat, "TenderActivityId": aid,
                    "Sort": "SubmitionDate"}
            first = etimad_json("/Tender/AllSupplierTendersForVisitorAsync", dict(base, PageNumber=1, SortDirection="DESC"))
            if not first:
                continue
            total = first.get("totalCount") or 0
            rows = list(first.get("data") or [])
            for p in range(2, min(3, math.ceil(total / 24)) + 1):
                d = etimad_json("/Tender/AllSupplierTendersForVisitorAsync", dict(base, PageNumber=p, SortDirection="DESC"))
                rows += (d or {}).get("data") or []
                time.sleep(0.4)
            if cat == 2 and total > 72:           # visitors only see 3 pages: take the oldest end as well
                for p in range(1, min(3, math.ceil((total - 72) / 24)) + 1):
                    d = etimad_json("/Tender/AllSupplierTendersForVisitorAsync", dict(base, PageNumber=p, SortDirection="ASC"))
                    rows += (d or {}).get("data") or []
                    time.sleep(0.4)
            for r in rows:
                et_add(store, r, status, name)
                n += 1
            progress(f"السعودية - اعتماد: {name} ({'مترسية' if cat == 6 else 'مطروحة'}) {len(rows)} من {total}")
            time.sleep(0.4)
    return n


# ----------------------------------------------------------------- source: UAE - Dubai eSupply
ESUPPLY = "https://esupply.dubai.gov.ae"


def collect_esupply(store, progress, max_pages=10):
    import http.cookiejar
    op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
    hdr = {"User-Agent": ct.UA["User-Agent"], "Accept": "text/html,*/*", "Accept-Language": "en"}
    list_url = ESUPPLY + "/esop/guest/go/public/opportunity/current"
    h = op.open(urllib.request.Request(list_url, headers=hdr), timeout=90).read().decode("utf-8", "replace")
    form = h[h.find('id="OpportunityListManager"'):]
    form = form[:form.find("</form>")]
    fields = {}
    for m in re.finditer(r"<input[^>]*>", form):
        t = m.group(0)
        nm, v = re.search(r'name="([^"]+)"', t), re.search(r'value="([^"]*)"', t)
        if nm and 'type="checkbox"' not in t:
            fields[nm.group(1)] = html.unescape(v.group(1)) if v else ""
    n, page = 0, 1
    while True:
        rows = re.findall(r'<tr class="table_cnt_body_[ab]">(.*?)</tr>', h, re.S)
        for tr in rows:
            tds = [ct.strip_tags(x) for x in re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S)]
            oid = re.search(r"goToDetail\(&#39;(\d+)&#39;", tr)
            if len(tds) < 6 or not oid:
                continue
            _cur, org, title, pub, cat, closing = tds[:6]
            clo = iso(any_date(closing))
            store.add("esupply", f"eSupply-{oid.group(1)}", org, title,
                      "مطروحة" if not clo or clo >= iso(TODAY()) else "مقفولة - تحت الدراسة",
                      official=(9, f"{ESUPPLY}/esop/toolkit/opportunity/current/{oid.group(1)}/detail.si", "دبي eSupply"),
                      url=list_url, country="AE", kind="مناقصة", publish=iso(any_date(pub)), closing=clo,
                      ttype=cat, notes="دبي - eSupply",
                      event={"date": iso(any_date(pub)), "etype": "طرح (دبي eSupply)", "url": list_url,
                             "text": f"{org} - نشر {pub} - إقفال {closing} - التصنيف {cat}"})
            n += 1
        if len(rows) < 100 or page >= max_pages:
            break
        page += 1
        fields.update({"opportunityListManager.pagerComponent.page": str(page), "userAct": "changePage"})
        h = op.open(urllib.request.Request(ESUPPLY + "/esop/toolkit/opportunity/current/list.si",
                                           data=urllib.parse.urlencode(fields).encode(),
                                           headers=dict(hdr, **{"Content-Type": "application/x-www-form-urlencoded"})),
                    timeout=90).read().decode("utf-8", "replace")
    progress(f"الإمارات - دبي eSupply: {n}")
    return n


# ----------------------------------------------------------------- source: news (Google News RSS)
NEWS = {
    "KW": {"gl": "KW", "queries": [
        "مناقصة الكويت", "طرح مناقصة الكويت", "ترسية مناقصة الكويت", "الجهاز المركزي للمناقصات العامة",
        "مشروع الكويت مليون دينار", "مشاريع كبرى الكويت", "خطة التنمية الكويت مشروع", "وزارة الأشغال العامة مشروع",
        "الرعاية السكنية مشروع مناقصة", "وزارة الكهرباء والماء مناقصة", "بلدية الكويت مشروع", "الهيئة العامة للطرق مشروع",
        "نفط الكويت عقد مشروع", "البترول الوطنية مناقصة", "الكويتية للصناعات البترولية المتكاملة مشروع",
        "الشراكة بين القطاعين العام والخاص الكويت مشروع", "ميناء مبارك الكبير", "مدينة جنوب سعد العبدالله",
        "مدينة المطلاع السكنية مشروع", "صيانة وتشغيل مناقصة الكويت", "Kuwait tender awarded", "Kuwait project contract"],
        "rx": r"كويت|kuwait",
        "outlets": r"القبس|الراي|الأنباء الكويتية|جريدة الأنباء$|الجريدة الكويتية|جريدة الجريدة|كونا|الوطن الكويتية|"
                   r"السياسة الكويتية|النهار الكويتية|Arab Times|Times Kuwait|KUNA"},
    "SA": {"gl": "SA", "queries": [
        "منافسة حكومية السعودية", "ترسية مشروع السعودية", "عقد تشغيل وصيانة السعودية", "مشاريع السعودية مليار ريال",
        "طرح منافسة وزارة", "وزارة البلديات والإسكان مشروع", "أرامكو عقد مشروع", "نيوم عقد مشروع",
        "الهيئة الملكية للجبيل وينبع مشروع", "الشركة السعودية للكهرباء عقد", "المياه الوطنية مشروع عقد",
        "تشغيل وصيانة مستشفيات السعودية عقد", "إدارة المرافق السعودية عقد", "مدن الصناعية مشروع عقد",
        "Saudi Arabia contract awarded", "Saudi facility management contract", "Saudi project tender"],
        "rx": r"سعودي|السعودية|المملكة|الرياض|جدة|الدمام|مكة|المدينة المنورة|نيوم|أرامكو|ارامكو|saudi|riyadh|jeddah|neom|aramco|ksa",
        "outlets": r"عكاظ|جريدة الرياض|الوطن السعودية|سبق|الاقتصادية|مال|أرقام|ارقام|Arab News|Saudi Gazette|Argaam|واس"},
    "AE": {"gl": "AE", "queries": [
        "مناقصة الإمارات", "مشروع دبي عقد", "أبوظبي مشروع عقد", "هيئة كهرباء ومياه دبي مشروع", "بلدية دبي مشروع",
        "هيئة الطرق والمواصلات دبي عقد", "الشارقة مشروع عقد", "UAE tender", "Dubai contract awarded",
        "Abu Dhabi project contract awarded", "UAE facility management contract"],
        "rx": r"الإمارات|الامارات|دبي|أبوظبي|ابوظبي|أبو ظبي|الشارقة|عجمان|رأس الخيمة|الفجيرة|uae|emirates|dubai|abu dhabi|sharjah",
        "outlets": r"البيان|الخليج|الاتحاد|الإمارات اليوم|وام|Gulf News|Khaleej Times|The National|Emirates 247"},
}


# news about our own companies - saved with query "شركة: ..." and shown in the Fawaz / KJAC tab
COMPANY_NEWS = {"gl": "KW", "days": 365, "queries": [
    '"فواز للتجارة والخدمات الهندسية"', '"شركة فواز للتجارة"', '"فواز للتبريد"', '"شركة فواز" مناقصة',
    '"الكويتية اليابانية للتكييف"', '"الكويتية اليابانية" تكييف', '"Fawaz Trading" Kuwait', '"Fawaz Refrigeration"',
    '"KJAC" Kuwait', '"Kuwait Japan Air Conditioning"', '"الحساوي" الكويت', '"عائلة الحساوي"', '"Al-Hasawi" Kuwait'],
    # the full company name in quotes is precise enough even when the name is only inside the article
    "trust": {'"فواز للتجارة والخدمات الهندسية"', '"شركة فواز للتجارة"', '"الكويتية اليابانية للتكييف"', '"عائلة الحساوي"'},
    "rx": r"فواز|fawaz|kjac|الكويتي[ةه] الياباني[ةه]|kuwait japan|الحساوي|hasawi", "outlets": r"(?!)"}


def news_country_ok(country, title, summary, source):
    c = NEWS[country]
    return bool(re.search(c["rx"], title + " " + summary, re.I) or re.search(c["outlets"], source or "", re.I))


def collect_news(progress, countries, days=150):
    con = db()
    since = iso(TODAY() - datetime.timedelta(days=days))
    n = 0
    seen_at = datetime.datetime.now().isoformat(timespec="seconds")
    groups = [(cc, NEWS[cc], "") for cc in countries]
    if "KW" in countries:
        groups.append(("KW", COMPANY_NEWS, "شركة: "))
    for cc, cfg, tag in groups:
        for q in cfg["queries"]:
            ar = any("؀" <= ch <= "ۿ" for ch in q)
            gdays = cfg.get("days", days)
            gsince = iso(TODAY() - datetime.timedelta(days=gdays))
            url = ("https://news.google.com/rss/search?q=" + urllib.parse.quote(q + f" when:{gdays}d") +
                   (f"&hl=ar&gl={cfg['gl']}&ceid={cfg['gl']}:ar" if ar else f"&hl=en-US&gl={cfg['gl']}&ceid={cfg['gl']}:en"))
            try:
                root = ET.fromstring(ct.http_get(url, binary=True))
            except Exception as e:      # noqa: BLE001 - skip one failing query
                progress(f"الأخبار: فشل '{q}': {e}")
                continue
            for it in root.iter("item"):
                title = (it.findtext("title") or "").strip()
                link = (it.findtext("link") or "").strip()
                src = (it.findtext("source") or "").strip()
                try:
                    d = datetime.datetime.strptime((it.findtext("pubDate") or "")[:25].strip(), "%a, %d %b %Y %H:%M:%S").date()
                except ValueError:
                    d = None
                if not d or iso(d) < gsince:
                    continue
                if src and title.endswith(" - " + src):
                    title = title[: -len(src) - 3]
                desc = ct.strip_tags(it.findtext("description") or "")
                if not (q in cfg.get("trust", ()) or re.search(cfg["rx"], title + " " + desc, re.I) or
                        re.search(cfg["outlets"], src or "", re.I)):
                    continue
                cur = con.execute("INSERT OR IGNORE INTO news(country,date,source,title,summary,url,query,cats,fawaz,seen_at) "
                                  "VALUES(?,?,?,?,?,?,?,?,?,?)",
                                  (cc, iso(d), src, title, desc[:400], link, tag + q, ",".join(categorize(title)),
                                   int(fawaz_scope(title) or bool(tag)), seen_at))
                if tag and not cur.rowcount:    # already there from a general search -> it is about us too
                    con.execute("UPDATE news SET query=?, fawaz=1 WHERE url=? AND query NOT LIKE 'شركة:%'", (tag + q, link))
                n += cur.rowcount
            con.commit()
            progress(f"الأخبار ({'فواز و KJAC' if tag else COUNTRIES[cc]}): {q}")
            time.sleep(0.3)
    con.close()
    return n


# ----------------------------------------------------------------- update (tracking)
SRC_NAME = {"capt": "الجهاز المركزي (المطروح)", "capt-minutes": "قرارات الجهاز (الكويت اليوم)", "capt-closing": "فض العطاءات (الجهاز المركزي)",
            "gazette": "جريدة الكويت اليوم", "pahw": "الرعاية السكنية", "adham": "تقرير مشاريع",
            "awards": "كشف الترسيات", "etimad": "منصة اعتماد", "esupply": "دبي eSupply"}


def ukey(it):
    if it["numkey"]:
        return f"{it['country']}|{it['numkey']}|{'-'.join(sorted(org_tokens(it['org'])))[:60]}"
    return f"{it['country']}|subj|{norm(it['subject'])[:90]}"


def is_fz(it):
    if fawaz_scope(it["subject"]):
        return True
    return "التشغيل والصيانة والنظافة" in (it["notes"] or "") and bool(re.search(r"صيانه|تشغيل", norm(it["subject"])))


def save_store(store):
    """upsert every merged item; log what changed since the last update"""
    now = datetime.datetime.now().isoformat(timespec="seconds")
    today = iso(TODAY())
    con = db()
    new_n = chg_n = 0
    # the very first load of a country is the baseline - only what appears after it counts as "new"
    populated = {r[0] for r in con.execute("SELECT DISTINCT country FROM items")}
    for it in store.items:
        it["sources"] = {SRC_NAME.get(s, s) for s in it["sources"]}
        for e in it["ev"]:
            e["source"] = SRC_NAME.get(e["source"], e["source"])
        if it["status"] in ("مطروحة", "") and it["closing"] and it["closing"] < today:
            it["status"] = "مقفولة - تحت الدراسة"
        if not it["status"] and it["sources"] == {SRC_NAME["capt-minutes"]}:
            it["status"] = "عقد جاري (قرار الجهاز)"     # extensions / change orders of running contracts
        if not it["status"]:
            it["status"] = "مطروحة" if it["closing"] and it["closing"] >= today else "مقفولة - تحت الدراسة"
        k = ukey(it)
        notes = (it["notes"] + " | المصادر: " + "، ".join(sorted(it["sources"]))).strip(" |")
        vals = dict(country=it["country"], kind=it["kind"], number=it["number"], numkey=it["numkey"], org=it["org"],
                    subject=it["subject"], status=it["status"], publish=it["publish"], closing=it["closing"], ptm=it["ptm"],
                    fees=it["fees"], bond=it["bond"], ttype=it["ttype"], value=it["value"], winner=it["winner"],
                    award_date=it["award_date"], cats=",".join(categorize(it["subject"], it["org"])), fawaz=int(is_fz(it)),
                    notes=notes, url=it["url"], official=it.get("official", (0, "", ""))[1],
                    official_label=it.get("official", (0, "", ""))[2],
                    docs=json.dumps(it.get("docs") or [], ensure_ascii=False) if it.get("docs") else "")
        old = con.execute("SELECT * FROM items WHERE ukey=?", (k,)).fetchone()
        if old is not None and not vals["docs"] and old["docs"]:
            vals["docs"] = old["docs"]
        if old is not None and not vals["official"] and old["official"]:
            vals["official"], vals["official_label"] = old["official"], old["official_label"]
        if old is None:
            cur = con.execute("INSERT INTO items(ukey," + ",".join(vals) + ",first_seen,last_seen,changed,updated) VALUES(" +
                              ",".join("?" * (len(vals) + 5)) + ")",
                              (k, *vals.values(), now if it["country"] in populated else BASELINE, now, "", now))
            iid = cur.lastrowid
            new_n += it["country"] in populated
        else:
            iid = old["id"]
            changed = False
            for f, label in TRACKED.items():
                ov, nv = old[f] or "", vals[f] or ""
                if f == "status" and ov in ("مترسية", "ملغاة") and STATUS_RANK.get(nv, -1) < STATUS_RANK.get(ov, -1):
                    vals[f] = ov                   # an award / cancellation is final
                    continue
                if f == "closing" and ov and nv and nv < ov:
                    vals[f] = ov                   # postponements only move the date forward
                    continue
                if nv and nv != ov:
                    con.execute("INSERT INTO changes(item_id,at,field,old,new) VALUES(?,?,?,?,?)", (iid, now, label, ov, nv))
                    changed = True
                elif not nv and ov:
                    vals[f] = ov                   # a source that knows less must not erase what we had
            if changed:
                chg_n += 1
            con.execute("UPDATE items SET " + ",".join(f"{c}=?" for c in vals) + ",last_seen=?,updated=?" +
                        (",changed=?" if changed else "") + " WHERE id=?",
                        (*vals.values(), now, now, *((now,) if changed else ()), iid))
        for e in it["ev"]:
            con.execute("INSERT OR IGNORE INTO events(item_id,date,source,etype,text,url,file,page) VALUES(?,?,?,?,?,?,?,?)",
                        (iid, e.get("date", ""), e["source"], e.get("etype", ""), e.get("text", ""), e.get("url", ""),
                         e.get("file", ""), e.get("page", 0)))
    # items we already had whose closing date passed since the last update
    for r in con.execute("SELECT id FROM items WHERE status='مطروحة' AND closing<>'' AND closing<?", (today,)).fetchall():
        con.execute("INSERT INTO changes(item_id,at,field,old,new) VALUES(?,?,?,?,?)", (r[0], now, "الحالة", "مطروحة", "مقفولة - تحت الدراسة"))
        con.execute("UPDATE items SET status='مقفولة - تحت الدراسة', changed=? WHERE id=?", (now, r[0]))
    con.commit()
    con.close()
    return new_n, chg_n


SOURCES = {"KW": [("ترسيات جديدة من الموقع", refresh_awards), ("الترسيات", collect_awards), ("جريدة الكويت اليوم", collect_gazette), ("تقرير مشاريع", collect_adham),
                  ("الجهاز المركزي", collect_capt), ("فض العطاءات", collect_capt_closing), ("الرعاية السكنية", collect_pahw)],
           "SA": [("منصة اعتماد", collect_etimad)],
           "AE": [("دبي eSupply", collect_esupply)]}


def run_update(progress=print, countries=None):
    countries = countries or list(COUNTRIES)
    started = datetime.datetime.now()
    store, summary = Store(), {}
    for cc in countries:
        for name, fn in SOURCES[cc]:
            try:
                summary[f"{COUNTRIES[cc]} - {name}"] = fn(store, progress)
            except Exception as e:      # noqa: BLE001 - one source failing must not stop the rest
                summary[f"{COUNTRIES[cc]} - {name}"] = f"خطأ: {e}"
                log(f"خطأ في {name}: {e}", progress)
    try:
        summary["أخبار جديدة"] = collect_news(progress, countries)
    except Exception as e:              # noqa: BLE001
        summary["أخبار جديدة"] = f"خطأ: {e}"
    new_n, chg_n = save_store(store)
    summary["جديد"], summary["اتغيّر"] = new_n, chg_n
    try:
        n_docs = prefetch_capt_docs(progress)
        if n_docs:
            summary["ملفات الجهاز المركزي"] = n_docs
        n_arch = archive_capt_docs(progress)
        if n_arch:
            summary["أرشيف الكراسات"] = n_arch
    except Exception as e:                  # noqa: BLE001
        summary["ملفات الجهاز المركزي"] = f"خطأ: {e}"
    if cfg.S.get("fawaz_watch", True):      # our own tenders + Sanjay's e-mails (PC only, confidential)
        try:
            import fawaz_watch
            summary["متابعة فواز"] = fawaz_watch.run(sys.modules[__name__], ct, progress)
        except Exception as e:              # noqa: BLE001
            summary["متابعة فواز"] = f"خطأ: {e}"
    try:                                    # phone app data (GitHub Pages) - only where mobile_publish is on
        import mobile_export
        if mobile_export.publish(progress):
            summary["نسخة التليفون"] = "اتحدثت"
    except Exception as e:                  # noqa: BLE001
        summary["نسخة التليفون"] = f"خطأ: {e}"
    con = db()
    con.execute("INSERT INTO runs(started,finished,summary) VALUES(?,?,?)",
                (started.isoformat(timespec="seconds"), datetime.datetime.now().isoformat(timespec="seconds"),
                 json.dumps(summary, ensure_ascii=False)))
    con.commit()
    con.close()
    log(f"التحديث خلص: {json.dumps(summary, ensure_ascii=False)}", progress)
    return summary


# ----------------------------------------------------------------- search
def all_items():
    con = db()
    rows = [dict(r) for r in con.execute("SELECT * FROM items")]
    con.close()
    return rows


_IDF = {}


def idf(rows):
    if _IDF.get("n") != len(rows):
        df = collections.Counter()
        for r in rows:
            df.update(set(stems(r["subject"] + " " + r["org"])))
        _IDF.clear()
        _IDF.update({k: math.log((len(rows) + 1) / (v + 0.5)) for k, v in df.items()})
        _IDF["n"] = len(rows)
    return _IDF


def search(q="", status="", cat="", org="", fawaz=False, closing_days=0, limit=250, country="", fresh=0):
    rows = all_items()
    tb = idf(rows)
    qs = [s for s in dict.fromkeys(stems(q)) if not s.isdigit()]
    qn = re.findall(r"\d+", norm(q))
    wsum = sum(tb.get(s, 8.0) for s in qs) or 1.0
    today = iso(TODAY())
    lim = iso(TODAY() + datetime.timedelta(days=closing_days)) if closing_days else ""
    since = (datetime.datetime.now() - datetime.timedelta(days=fresh)).isoformat(timespec="seconds") if fresh else ""
    out = []
    for r in rows:
        if country and r["country"] != country:
            continue
        if since and not (r["first_seen"] >= since or (r["changed"] or "") >= since):
            continue
        if status and r["status"] != status:
            continue
        if cat and cat not in (r["cats"] or ""):
            continue
        if org and org_sim(org, r["org"]) < 0.6:
            continue
        if fawaz and not r["fawaz"]:
            continue
        if lim and not (r["closing"] and today <= r["closing"] <= lim):
            continue
        sc = 1.0
        if qs or qn:
            hay = set(stems(r["subject"] + " " + r["org"] + " " + (r["winner"] or "")))
            got = sum(tb.get(s, 8.0) * (1 if s in hay else 0.6 if len(s) >= 4 and any(h.startswith(s) for h in hay) else 0) for s in qs)
            sc = got / wsum if qs else 0
            if qn:
                ns = 1.0 if all(x in ct.num_parts(r["number"]) for x in qn) else 0
                sc = max(sc, ns) if not qs else 0.65 * sc + 0.35 * ns
            if sc < 0.4:
                continue
        out.append((sc, r))
    if q:
        out.sort(key=lambda t: (round(t[0], 2), t[1]["closing"] or t[1]["award_date"] or ""), reverse=True)
    else:
        rank = {"مطروحة": 0, "اتوافق على طرحها": 1, "مقفولة - تحت الدراسة": 2, "أسعار معلنة": 3, "مترسية": 4, "ملغاة": 5}
        out.sort(key=lambda t: (rank.get(t[1]["status"], 9), t[1]["closing"] or "9999" if t[1]["status"] == "مطروحة"
                                else "~" + (t[1]["award_date"] or t[1]["publish"] or "")))
    return [dict(r, score=round(s, 3)) for s, r in out[:limit]]


def search_news(q, limit=80, fawaz=False, country=""):
    con = db()
    rows = [dict(r) for r in con.execute("SELECT * FROM news WHERE (?='' OR country=?) ORDER BY date DESC", (country, country))]
    con.close()
    if fawaz:
        rows = [r for r in rows if r["fawaz"]]
    if not q:
        return rows[:limit]
    qs = set(s for s in stems(q) if len(s) > 2)
    out = []
    for r in rows:
        hay = set(stems(r["title"] + " " + (r["summary"] or "")))
        sc = len(qs & hay) / len(qs) if qs else 0
        if sc >= 0.5:
            out.append((sc, r))
    out.sort(key=lambda t: (round(t[0], 1), t[1]["date"]), reverse=True)
    return [dict(r, score=round(s, 2)) for s, r in out[:limit]]


# ----------------------------------------------------------------- unread counts (badges on the tabs)
US_RX = r"فواز|fawaz|kjac|الكويتي[ةه] الياباني[ةه]|kuwait japan"


def _seen_store():
    import fawaz_watch
    return fawaz_watch


def tab_markers():
    """the current 'everything up to here was read' point of each tab"""
    con = db()
    m = {"items": datetime.datetime.now().isoformat(timespec="seconds"),
         "news": datetime.datetime.now().isoformat(timespec="seconds")}
    con.close()
    try:
        m["awards"] = str(len(ct.read_excel_rows()))
    except Exception:                   # noqa: BLE001
        m["awards"] = "0"
    fw = _seen_store()
    if fw.DB.exists():
        c = fw.db()
        m["watch"] = str(c.execute("SELECT coalesce(max(id),0) FROM feed").fetchone()[0])
        c.close()
    else:
        m["watch"] = "0"
    return m


def unread():
    """-> ({tab: count}, {tab: marker}); a tab never opened starts at 'nothing unread'"""
    fw = _seen_store()
    seen = fw.get_seen()
    cur = None
    for tab in ("items", "news", "awards", "watch"):
        if tab not in seen:
            cur = cur or tab_markers()
            fw.set_seen(tab, cur[tab])
            seen[tab] = cur[tab]
    con = db()
    n = {"items": con.execute("SELECT count(*) FROM items WHERE (first_seen>? AND first_seen>?) OR changed>?",
                              (seen["items"], BASELINE, seen["items"])).fetchone()[0],
         "news": con.execute("SELECT count(*) FROM news WHERE seen_at>? AND query NOT LIKE 'شركة:%'", (seen["news"],)).fetchone()[0]}
    con.close()
    try:
        n["awards"] = max(0, len(ct.read_excel_rows()) - int(seen["awards"] or 0))
    except Exception:                   # noqa: BLE001
        n["awards"] = 0
    n["watch"] = fw.unread_feed(seen["watch"])
    return n, seen


def mark_seen(tab):
    m = tab_markers()
    if tab in m:
        _seen_store().set_seen(tab, m[tab])


def watch_view():
    """everything for the Fawaz / KJAC tab"""
    fw = _seen_store()
    d = fw.data()
    con = db()
    d["company_news"] = [dict(r) for r in con.execute(
        "SELECT id, date, source, title, url, seen_at FROM news WHERE query LIKE 'شركة:%' ORDER BY date DESC LIMIT 150")]
    d["won_items"] = [dict(r) for r in con.execute(
        "SELECT id, country, number, org, subject, status, winner, value, award_date, official FROM items "
        "WHERE winner<>'' ORDER BY award_date DESC")]
    con.close()
    d["won_items"] = [r for r in d["won_items"] if re.search(US_RX, r["winner"], re.I)][:150]
    import gazette_online as go
    d["gazette"] = dict(gazette_status(), site=GZ_SITE, folder=str(GAZETTE_PDF_DIR), account=go.account()[0], has_account=all(go.account()))
    try:
        d["awards"] = [r for r in ct.read_excel_rows() if re.search(US_RX, r["winner"] or "", re.I)][:200]
    except Exception:                   # noqa: BLE001
        d["awards"] = []
    return d


def accounts_api(body, log):
    """add / reset / (de)activate / Fawaz permission / delete / mail -> the list; any change republishes"""
    import accounts, phone_lock, mobile_export as me
    owner = accounts.ensure_owner()
    act, user, out, changed = body.get("action", "list"), body.get("u", ""), {"msg": ""}, False
    try:
        if act == "add":
            u, pw = accounts.add_user(body.get("name", ""), note=body.get("note", ""), perms=body.get("perms"),
                                      countries=body.get("countries"), expires=_until(body.get("months")))
            out.update(new={"u": u, "pw": pw}, msg=f"اتعمل حساب {u}")
            changed = True
        elif act == "extend":                       # +N months from today or from the current end, whichever is later
            d = accounts.load()["users"][accounts.norm_name(user)]
            months = int(body.get("months") or 0)
            if months <= 0:
                accounts.update(user, expires="")
                out["msg"] = "بقى من غير مدة"
            else:
                base = max(datetime.date.today(), datetime.date.fromisoformat(d["expires"])) if d.get("expires") else datetime.date.today()
                accounts.update(user, expires=_until(months, base), active=True)
                out["msg"] = f"الاشتراك بقى لحد {_until(months, base)}"
            changed = True
        elif act == "perms":
            accounts.update(user, perms=body.get("perms") or [], countries=body.get("countries") or [])
            out["msg"], changed = "الصلاحيات اتحفظت", True
        elif act == "reset":
            pw = accounts.set_password(user)
            out.update(new={"u": accounts.norm_name(user), "pw": pw}, msg="اتعملت كلمة سر جديدة")
            changed = True
        elif act in ("active", "fawaz"):
            if act == "active" and accounts.norm_name(user) == owner and not body.get("value"):
                raise ValueError("ده حسابك - مينفعش يتوقف")
            accounts.update(user, **{act: bool(body.get("value"))})
            out["msg"] = {("active", True): "اتفعّل", ("active", False): "اتوقف - مش هيقدر يفتح من أول تحديث",
                          ("fawaz", True): "بقى يشوف بيانات فواز السرية", ("fawaz", False): "مبقاش يشوف بيانات فواز السرية"
                          }[(act, bool(body.get("value")))]
            changed = True
        elif act == "delete":
            accounts.delete(user)
            out["msg"], changed = "اتمسح", True
        elif act == "mail":                         # send a new user's login to the owner's Gmail
            err = phone_lock.mail_password(accounts.norm_name(user), body.get("pw", ""), "بيانات دخول تطبيق الموبايل:")
            out["msg"] = err or f"اتبعتت على {phone_lock.owner_email()}"
    except (ValueError, KeyError) as e:
        out["msg"] = f"خطأ: {e}"
    if changed:                                     # the phone site is re-encrypted for the new list right away
        threading.Thread(target=me.publish, args=(log, True), daemon=True).start()
    out.update(users=accounts.listing(), owner=owner, owner_password=cfg.S.get("phone_owner_password", ""),
               email=phone_lock.owner_email(), site=me.PAGES, features=accounts.FEATURES, countries=accounts.COUNTRIES,
               today=datetime.date.today().isoformat())
    return out


def _until(months, base=None):
    """end date N months after base (today) - '' = no end"""
    months = int(months or 0)
    if months <= 0:
        return ""
    base = base or datetime.date.today()
    y, m = divmod(base.month - 1 + months, 12)
    import calendar
    day = min(base.day, calendar.monthrange(base.year + y, m + 1)[1])
    return datetime.date(base.year + y, m + 1, day).isoformat()



# ----------------------------------------------------------------- tender documents
DOCS_OUT = BASE / "مستندات المناقصات"


def _capt_q(official, number=""):
    """(ministry code, tender number exactly as CAPT writes it) from the item's official CAPT link"""
    q = urllib.parse.parse_qs(urllib.parse.urlparse(official or "").query)
    return q.get("ministry_code", [""])[0], q.get("tender_no", [number])[0] or number


def item_docs(iid, live=True):
    """every document we know for an item: the source's files (PAHW), CAPT's files (with the owner's
    account), CAPT minutes, gazette pages -> {"item", "docs": [{label,url,date,src,local}], "capt": {...}}"""
    import tender_docs as td
    con = db()
    it = dict(con.execute("SELECT * FROM items WHERE id=?", (iid,)).fetchone())
    ev = [dict(e) for e in con.execute("SELECT date, source, etype, text, url, file, page FROM events WHERE item_id=?", (iid,))]
    cached = con.execute("SELECT docs, fetched, logged FROM capt_docs WHERE item_id=?", (iid,)).fetchone()
    con.close()
    docs = [dict(d, local=False) for d in json.loads(it.get("docs") or "[]")]
    capt = None
    off = it.get("official") or ""
    if "capt.gov.kw/ar/tenders/opening-tenders" in off and it.get("number"):
        if live:
            try:
                code, tno = _capt_q(off, it["number"])
                r = td.capt_docs(tno, code)
                kept = json.loads(cached["docs"]) if cached else []
                if not any(d.get("url") for d in r.get("docs", [])) and any(d.get("url") for d in kept):
                    r = {"docs": kept, "login": bool(cached["logged"]), "error": "",
                         "note": "المناقصة اتشالت من صفحة الجهاز - دي اللينكات اللي البرنامج سجّلها وهي مطروحة"}
                else:
                    save_capt_docs(iid, r)
            except Exception as e:                  # noqa: BLE001
                r = {"docs": json.loads(cached["docs"]) if cached else [], "login": bool(cached and cached["logged"]),
                     "error": f"موقع الجهاز ما ردّش: {e}"}
        else:
            r = {"docs": json.loads(cached["docs"]) if cached else [], "login": bool(cached and cached["logged"]), "error": ""}
        capt = {"login": r.get("login"), "error": r.get("error", ""), "account": td.account()[0], "note": r.get("note", "")}
        docs += [dict(d, src="الجهاز المركزي", local=False) for d in r.get("docs", [])]
    for e in ev:                                    # CAPT minutes (public PDFs) + gazette pages on this PC
        if (e["url"] or "").lower().endswith(".pdf") and "capt.gov.kw" in (e["url"] or ""):
            docs.append({"label": e["etype"] or "محضر الجهاز", "url": e["url"], "date": e["date"], "src": "محاضر الجهاز", "local": False})
        if e["file"] and str(e["file"]).lower().endswith(".pdf") and e["page"]:
            docs.append({"label": f"صفحة الجريدة {e['page']}", "url": f"/api/file?path={urllib.parse.quote(e['file'])}#page={e['page']}",
                         "date": e["date"], "src": "الكويت اليوم", "local": True, "file": e["file"], "page": e["page"]})
    if off.lower().endswith(".pdf"):
        docs.append({"label": it.get("official_label") or "المحضر", "url": off, "date": "", "src": "رسمي", "local": False})
    for f in archived_files(iid):                   # kept from when the tender was still published
        docs.append({"label": f.stem, "url": "", "path": str(f), "date": "", "src": "أرشيف البرنامج", "local": True, "archived": True})
    seen, out = set(), []
    for d in docs:
        k = d.get("url") or d.get("label")
        if k not in seen:
            seen.add(k)
            out.append(d)
    return {"item": {k: it.get(k) for k in ("id", "number", "subject", "org", "country", "official")}, "docs": out, "capt": capt}


def save_capt_docs(iid, r):
    con = db()
    old = con.execute("SELECT docs FROM capt_docs WHERE item_id=?", (iid,)).fetchone()
    if old and any(d.get("url") for d in json.loads(old["docs"] or "[]")) and not any(d.get("url") for d in r.get("docs", [])):
        con.close()                                 # never replace links we have with nothing
        return
    con.execute("INSERT INTO capt_docs(item_id, docs, fetched, logged) VALUES(?,?,?,?) ON CONFLICT(item_id) DO UPDATE SET "
                "docs=excluded.docs, fetched=excluded.fetched, logged=excluded.logged",
                (iid, json.dumps(r.get("docs", []), ensure_ascii=False), datetime.datetime.now().isoformat(timespec="seconds"),
                 int(bool(r.get("login")))))
    con.commit()
    con.close()


def download_docs(iid):
    """download every reachable document of an item into its own folder (gazette: just the page)"""
    import tender_docs as td
    d = item_docs(iid)
    files = [x for x in d["docs"] if x.get("url") and not x.get("local")]
    folder, res = td.download(d["item"], files, DOCS_OUT)
    for x in d["docs"]:                              # gazette pages -> one-page PDFs
        if x.get("local") and x.get("file"):
            try:
                import pymupdf
                src = pymupdf.open(x["file"])
                out = pymupdf.open()
                out.insert_pdf(src, from_page=int(x["page"]) - 1, to_page=int(x["page"]) - 1)
                name = td.safe(f"الكويت اليوم ص {x['page']} - {Path(x['file']).stem}", 80) + ".pdf"
                out.save(folder / name)
                res.append((name, True, "صفحة الجريدة"))
            except Exception as e:                   # noqa: BLE001
                res.append((x["label"], False, str(e)[:100]))
    locked = [x["label"] for x in d["docs"] if not x.get("url")]
    return {"folder": str(folder), "files": res, "locked": locked, "capt": d["capt"]}


def prefetch_capt_docs(progress, limit=40):
    """with the owner's CAPT account: keep the file lists of open CAPT tenders fresh (a few per update),
    so the phone can show them too"""
    import tender_docs as td
    if not all(td.account()):
        return 0
    old = (datetime.datetime.now() - datetime.timedelta(hours=20)).isoformat(timespec="seconds")
    con = db()
    rows = con.execute("SELECT i.id, i.number, i.official FROM items i LEFT JOIN capt_docs c ON c.item_id=i.id "
                       "WHERE i.country='KW' AND i.status='مطروحة' AND i.official LIKE '%capt.gov.kw/ar/tenders/opening-tenders%' "
                       "AND (c.fetched IS NULL OR c.fetched<? OR c.logged=0) ORDER BY i.closing LIMIT ?", (old, limit)).fetchall()
    con.close()
    n = 0
    for r in rows:
        try:
            code, tno = _capt_q(r["official"], r["number"])
            save_capt_docs(r["id"], td.capt_docs(tno, code))
            n += 1
        except Exception as e:                       # noqa: BLE001
            progress(f"ملفات الجهاز المركزي: {e}")
            break
        time.sleep(0.5)
    return n


def archive_capt_docs(progress, limit=6):
    """download the CAPT files of open tenders (Fawaz scope by default) while they are still published -
    after closing / award the site removes them, the archive keeps them"""
    import tender_docs as td
    scope = cfg.S.get("archive_scope", "fawaz")     # fawaz / all / off
    if scope == "off" or not all(td.account()):
        return 0
    con = db()
    rows = con.execute("SELECT i.id, i.number, i.subject, i.fawaz, c.docs FROM items i JOIN capt_docs c ON c.item_id=i.id "
                       "WHERE c.logged=1 AND (c.archived IS NULL OR c.archived='') AND i.status='مطروحة'" +
                       (" AND i.fawaz=1" if scope == "fawaz" else "") + " LIMIT ?", (limit,)).fetchall()
    con.close()
    n = 0
    for r in rows:
        files = [d for d in json.loads(r["docs"] or "[]") if d.get("url")]
        if not files:
            continue
        folder, res = td.download({"number": r["number"], "subject": r["subject"]}, files, DOCS_OUT)
        if any(ok for _n, ok, _x in res):
            con = db()
            con.execute("UPDATE capt_docs SET archived=? WHERE item_id=?", (str(folder), r["id"]))
            con.commit()
            con.close()
            n += 1
    if n:
        progress(f"أرشيف الكراسات: اتحفظت ملفات {n} مناقصة")
    return n


def archived_files(iid):
    con = db()
    r = con.execute("SELECT archived FROM capt_docs WHERE item_id=?", (iid,)).fetchone()
    con.close()
    folder = Path(r["archived"]) if r and r["archived"] else None
    return sorted(folder.iterdir()) if folder and folder.is_dir() else []


def detail(iid):
    con = db()
    it = dict(con.execute("SELECT * FROM items WHERE id=?", (iid,)).fetchone())
    ev = [dict(e) for e in con.execute("SELECT * FROM events WHERE item_id=? ORDER BY date DESC", (iid,))]
    chg = [dict(c) for c in con.execute("SELECT at, field, old, new FROM changes WHERE item_id=? ORDER BY at DESC", (iid,))]
    con.close()
    hist, hits = [], []
    if it["country"] == "KW" and it["subject"]:       # awards file + CAPT minutes are Kuwait only
        rec = {"tno": it["number"], "subject": it["subject"], "org": it["org"], "row": -1}
        hist = ct.similar_history(rec)
        hits = ct.search_minutes(it["number"], it["subject"], it["org"], it["award_date"] or "", limit=15)
        if it["number"] and len(ct.num_parts(it["number"])) >= 2:
            hits = [h for h in hits if h["num"] > 0] or hits[:5]
    elif it["subject"]:                                # same country, similar work, already awarded
        tb = idf(all_items())
        qs = list(dict.fromkeys(stems(it["subject"])))
        wsum = sum(tb.get(s, 8.0) for s in qs) or 1.0
        for r in all_items():
            if r["country"] != it["country"] or r["status"] != "مترسية" or r["id"] == it["id"]:
                continue
            sc = sum(tb.get(s, 8.0) for s in qs if s in set(stems(r["subject"]))) / wsum
            if sc >= 0.5:
                hist.append({"date": r["award_date"] or r["publish"], "tno": r["number"], "subject": r["subject"],
                             "org": r["org"], "winner": r["winner"] or "-", "total": r["value"] or "-", "score": sc})
        hist = sorted(hist, key=lambda x: (round(x["score"], 1), x["date"]), reverse=True)[:12]
    news = search_news(it["subject"], limit=12, country=it["country"])
    bids = bid_sheet_for(it["number"], it["org"]) if it["country"] == "KW" and it["number"] else None
    return {"item": it, "events": ev, "changes": chg, "history": hist, "hits": hits, "news": news, "bids": bids}


# ----------------------------------------------------------------- PDF
def build_pdf(iid, hits=None, include_gazette=True):
    import pymupdf
    d = detail(iid)
    it, esc = d["item"], ct.esc
    hits = d["hits"] if hits is None else hits
    OUT.mkdir(exist_ok=True)
    fz = " class='fz'" if ct.is_fawaz(it["winner"]) else ""
    kv = [("النوع", it["kind"]), ("رقم المناقصة", it["number"] or "-"), ("الجهة", it["org"]), ("الموضوع", it["subject"]),
          ("الحالة", it["status"]), ("تاريخ الطرح", dmy(any_date(it["publish"]))), ("آخر موعد للعطاء", dmy(any_date(it["closing"]))),
          ("الاجتماع التمهيدي", it["ptm"]), ("سعر الوثائق", it["fees"]), ("التأمين الأولي", it["bond"]),
          ("نوع الطرح", it["ttype"]), ("الفائز", it["winner"]), ("القيمة", it["value"]),
          ("تاريخ الترسية", dmy(any_date(it["award_date"]))), ("التصنيف", it["cats"]), ("ملاحظات", it["notes"])]
    kv.insert(0, ("الدولة", COUNTRIES.get(it["country"], it["country"])))
    H = ["<h1>دليل مناقصات ومشاريع الخليج</h1>",
         f"<p class='small'>اتعمل {dmy(TODAY())} - شركة فواز للتجارة والخدمات الهندسية</p><h2>بيانات المناقصة</h2>"]
    H += [f"<p class='kv'><b>{k}:</b> <span{fz if k == 'الفائز' else ''}>{esc(v)}</span></p>" for k, v in kv if v]
    link = it.get("official") or it.get("url") or ""
    if link.startswith("http"):
        H.append(f"<p class='kv'><b>الإعلان الرسمي ({esc(it.get('official_label') or 'المصدر')}):</b> "
                 f"<a href='{esc(link)}'>{esc(link)[:95]}</a></p>")
    if d["changes"]:
        H.append("<h2>التغييرات اللي اتسجلت عليها</h2><table><tr><th>وقت الرصد</th><th>البند</th><th>كان</th><th>بقى</th></tr>")
        H += [f"<tr><td>{esc(c['at'][:16].replace('T', ' '))}</td><td>{esc(c['field'])}</td><td>{esc(c['old'])}</td><td>{esc(c['new'])}</td></tr>" for c in d["changes"]]
        H.append("</table>")
    if d["events"]:
        H.append("<h2>كل اللي اتنشر عنها (الأحدث الأول)</h2><table><tr><th>التاريخ</th><th>المصدر</th><th>الحدث</th><th>التفاصيل</th></tr>")
        H += [f"<tr><td>{dmy(any_date(e['date']))}</td><td>{esc(e['source'])}</td><td>{esc(e['etype'])}</td><td>{esc(e['text'])[:260]}</td></tr>" for e in d["events"]]
        H.append("</table>")
    if d["history"]:
        H.append("<h2>ترسيات سابقة لنفس نوع الشغل (للتسعير)</h2><table><tr><th>التاريخ</th><th>الرقم</th><th>الموضوع</th><th>الجهة</th><th>الفائز</th><th>القيمة</th></tr>")
        for x in d["history"]:
            f2 = " class='fz'" if ct.is_fawaz(x["winner"]) else ""
            H.append(f"<tr><td>{dmy(any_date(x['date']))}</td><td>{esc(x['tno'])}</td><td>{esc(x['subject'])[:120]}</td><td>{esc(x['org'])}</td><td{f2}>{esc(x['winner'])[:90]}</td><td>{esc(x['total'])}</td></tr>")
        H.append("</table>")
    if d["news"]:
        H.append("<h2>أخبار قريبة من الموضوع</h2><table><tr><th>التاريخ</th><th>المصدر</th><th>العنوان</th></tr>")
        H += [f"<tr><td>{dmy(any_date(n['date']))}</td><td>{esc(n['source'])}</td><td><a href='{esc(n['url'])}'>{esc(n['title'])}</a></td></tr>" for n in d["news"]]
        H.append("</table>")
    if hits:
        H.append("<h2>صفحات محاضر الجهاز المركزي المرفقة</h2><table><tr><th>الاجتماع</th><th>التاريخ</th><th>النوع (تقريبي)</th><th>الملف / الصفحة</th></tr>")
        H += [f"<tr><td>{h['year']}/{h['meeting']}</td><td>{dmy(any_date(h['mdate']))}</td><td>{esc(h['type'])}</td><td>{esc(Path(h['file']).name)} ص {h['page']}</td></tr>" for h in sorted(hits, key=lambda h: h['mdate'], reverse=True)]
        H.append("</table>")
    out = ct.story_pages("".join(H))
    toc = [[1, "بيانات المناقصة", 1]]
    nums = [n for n in ct.num_parts(it["number"]) if not ct.is_year(n)]
    tb = idf(all_items())
    hl = set(sorted(set(stems(it["subject"])), key=lambda s: -tb.get(s, 8.0))[:6])
    added = set()
    for h in sorted(hits, key=lambda h: h["mdate"], reverse=True):
        k = (h["file"], h["page"])
        if k in added:
            continue
        added.add(k)
        src = pymupdf.open(ct.ARCHIVE / h["file"])
        out.insert_pdf(src, from_page=h["page"] - 1, to_page=h["page"] - 1)
        src.close()
        label = f"محضر {h['year']}/{h['meeting']} - {dmy(any_date(h['mdate']))} - {h['type'] or 'بند'} - {Path(h['file']).name} ص {h['page']}"
        ct.highlight(out[-1], nums, hl, ct.is_fawaz(it["winner"]))
        ct.stamp(out[-1], label)
        toc.append([1, label, out.page_count])
    if include_gazette:
        for e in d["events"]:
            if e["file"] and e["page"] and Path(e["file"]).exists() and (e["file"], e["page"]) not in added:
                added.add((e["file"], e["page"]))
                src = pymupdf.open(e["file"])
                if e["page"] <= src.page_count:
                    out.insert_pdf(src, from_page=e["page"] - 1, to_page=e["page"] - 1)
                    label = f"{e['source']} - {dmy(any_date(e['date']))} - {Path(e['file']).name} ص {e['page']}"
                    ct.stamp(out[-1], label)
                    toc.append([1, label, out.page_count])
                src.close()
    out.set_toc(toc)
    name = f"{ct.safe_name(it['number'] or it['kind'], 35)} - {ct.safe_name(it['subject'], 60)} - {TODAY():%Y-%m-%d}.pdf"
    path = OUT / name
    try:
        out.save(path, garbage=3, deflate=True)
    except Exception:                    # noqa: BLE001 - open in a viewer
        path = OUT / (path.stem + f" ({datetime.datetime.now():%H%M%S}).pdf")
        out.save(path, garbage=3, deflate=True)
    out.close()
    return path


# ----------------------------------------------------------------- web UI
def stats(country=""):
    con = db()
    w, a = ("WHERE country=?", (country,)) if country else ("", ())
    aw = (" AND country=?", (country,)) if country else ("", ())
    since = (datetime.datetime.now() - datetime.timedelta(days=3)).isoformat(timespec="seconds")
    runs = con.execute("SELECT * FROM runs ORDER BY id DESC LIMIT 1").fetchone()
    s = {"status": dict(con.execute(f"SELECT status, count(*) FROM items {w} GROUP BY status", a).fetchall()),
         "total": con.execute(f"SELECT count(*) FROM items {w}", a).fetchone()[0],
         "news": con.execute(f"SELECT count(*) FROM news {w}", a).fetchone()[0],
         "orgs": [r[0] for r in con.execute(f"SELECT org FROM items WHERE org<>''{aw[0]} GROUP BY org ORDER BY count(*) DESC LIMIT 150", aw[1])],
         "last_run": dict(runs) if runs else {},
         "closing_week": con.execute(f"SELECT count(*) FROM items WHERE status='مطروحة' AND closing<=?{aw[0]}",
                                     (iso(TODAY() + datetime.timedelta(days=7)), *aw[1])).fetchone()[0],
         "fresh": con.execute(f"SELECT count(*) FROM items WHERE (first_seen>=? OR changed>=?){aw[0]}", (since, since, *aw[1])).fetchone()[0],
         "countries": dict(con.execute("SELECT country, count(*) FROM items GROUP BY country").fetchall()),
         "country_names": COUNTRIES, "auto_minutes": auto_minutes(),
         "cats": [c for c, _ in CATS] + ["أخرى"]}
    con.close()
    return s


def last_run_age_hours():
    if not DB.exists():
        return 1e9
    con = db()
    r = con.execute("SELECT finished FROM runs ORDER BY id DESC LIMIT 1").fetchone()
    con.close()
    if not r or not r[0]:
        return 1e9
    return (datetime.datetime.now() - datetime.datetime.fromisoformat(r[0])).total_seconds() / 3600


def serve(open_browser=True):
    from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
    import webbrowser
    state = {"update": None, "log": [], "next": None}

    def start_update(reason=""):
        if state["update"] == "running":
            return
        state["update"], state["log"] = "running", ([reason] if reason else [])

        def job():
            try:
                s = run_update(lambda m: state["log"].append(m))
                state["log"].append("خلص التحديث: جديد " + str(s.get("جديد")) + " - اتغيّر " + str(s.get("اتغيّر")))
                state["update"] = "done"
            except Exception as e:      # noqa: BLE001
                state["log"].append(f"خطأ: {e}")
                state["update"] = "error"
        threading.Thread(target=job, daemon=True).start()

    def scheduler():
        # periodic update from the internet while the program runs; the interval can change from the window
        last_check = 0
        while True:
            if time.time() - last_check >= 1800:
                last_check = time.time()
                try:                     # installed program: new version on GitHub -> install it
                    if updater.check_and_apply(lambda m: state["log"].append(m)):
                        os._exit(0)
                except Exception:        # noqa: BLE001
                    pass
            if cfg.S.get("mobile_publish") and time.time() - state.get("phone_poll", 0) >= 60:
                state["phone_poll"] = time.time()      # "forgot the passphrase" requests from the phone
                try:
                    import phone_lock
                    threading.Thread(target=phone_lock.poll, args=(lambda m: state["log"].append(m),), daemon=True).start()
                except Exception:        # noqa: BLE001
                    pass
            every = auto_minutes()
            if not every:
                state["next"] = None
            elif state["update"] != "running":
                age = last_run_age_hours() * 60
                if age >= every:
                    start_update(f"تحديث تلقائي (آخر تحديث من {age:.0f} دقيقة)" if age < 1e8 else "أول تحديث")
                    age = 0
                state["next"] = (datetime.datetime.now() + datetime.timedelta(minutes=every - age)).isoformat(timespec="minutes")
            time.sleep(20)

    class H(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def send(self, obj, code=200, ctype="application/json; charset=utf-8"):
            data = obj if isinstance(obj, bytes) else json.dumps(obj, ensure_ascii=False, default=str).encode()
            self.send_response(code)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(data)))
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            self.wfile.write(data)

        def do_GET(self):
            u = urllib.parse.urlparse(self.path)
            q = {k: v[0] for k, v in urllib.parse.parse_qs(u.query).items()}
            try:
                if u.path == "/":
                    return self.send((cfg.res_dir() / "ui" / "واجهة الدليل.html").read_bytes(), ctype="text/html; charset=utf-8")
                if u.path == "/api/status":
                    try:
                        un, seen = unread()
                    except Exception:   # noqa: BLE001
                        un, seen = {}, {}
                    return self.send(dict(stats(q.get("country", "")), update=state["update"], log=state["log"][-18:],
                                          next=state["next"], version=cfg.build_info().get("version"),
                                          unread=un, seen=seen))
                if u.path == "/api/watch":
                    return self.send(watch_view())
                if u.path == "/api/gazette-account":
                    import gazette_online as go
                    if q.get("explore"):
                        return self.send(go.explore())
                    return self.send({"user": go.account()[0], "has_password": go.account()[1], "status": gazette_status()})
                if u.path == "/api/gazette/search":
                    return self.send(gazette_search(q.get("q", ""), issue=q.get("issue", "")))
                if u.path == "/api/docs":
                    return self.send(item_docs(int(q["id"])))
                if u.path == "/api/capt-account":
                    import tender_docs as td
                    user, has = td.account()
                    if q.get("explore"):                # what the company account area offers (menu links only)
                        return self.send(td.explore_account())
                    return self.send({"user": user, "has_password": has})
                if u.path == "/api/search":
                    return self.send(search(q.get("q", ""), q.get("status", ""), q.get("cat", ""), q.get("org", ""),
                                            q.get("fawaz") == "1", int(q.get("days") or 0), country=q.get("country", ""),
                                            fresh=int(q.get("fresh") or 0)))
                if u.path == "/api/news":
                    return self.send(search_news(q.get("q", ""), fawaz=q.get("fawaz") == "1", country=q.get("country", "")))
                if u.path == "/api/detail":
                    return self.send(detail(int(q["id"])))
                # --- Kuwait awards + minutes (the awards tool, same program)
                if u.path == "/api/awards/search":
                    return self.send(ct.search_awards(q.get("q", ""), limit=60, year=q.get("year") or None))
                if u.path == "/api/awards/detail":
                    rec = next(r for r in ct.read_excel_rows() if r["row"] == int(q["row"]))
                    return self.send({"rec": rec, "same": ct.same_tender(rec), "history": ct.similar_history(rec),
                                      "hits": ct.search_minutes(rec["tno"], rec["subject"], rec["org"], rec["date"])})
                if u.path == "/api/awards/minutes":
                    return self.send(ct.minutes_free_search(q.get("q", "")))
                if u.path == "/api/file":            # local official PDFs (gazette issues, minutes)
                    f = Path(q["path"]).resolve()
                    allowed = [GAZETTE_PDF_DIR.resolve(), ct.MOM.resolve()]
                    if f.suffix.lower() == ".pdf" and f.exists() and any(str(f).startswith(str(a)) for a in allowed):
                        return self.send(f.read_bytes(), ctype="application/pdf")
                    return self.send({"error": "file not allowed"}, 403)
                if u.path == "/api/page.png":
                    import pymupdf
                    f = Path(q["file"])
                    f = f if f.is_absolute() else ct.ARCHIVE / f
                    dd = pymupdf.open(f)
                    png = dd[int(q["page"]) - 1].get_pixmap(dpi=110).tobytes("png")
                    dd.close()
                    return self.send(png, ctype="image/png")
                self.send({"error": "not found"}, 404)
            except Exception as e:      # noqa: BLE001
                self.send({"error": str(e)}, 500)

        def do_POST(self):
            u = urllib.parse.urlparse(self.path)
            body = json.loads(self.rfile.read(int(self.headers.get("Content-Length", 0)) or 0) or b"{}")
            try:
                if u.path == "/api/report":
                    p = build_pdf(int(body["id"]), body.get("hits"))
                    os.startfile(p)
                    return self.send({"name": p.name})
                if u.path == "/api/awards/report":
                    rec = next((r for r in ct.read_excel_rows() if r["row"] == int(body["row"])), None) if body.get("row") else None
                    p, _ = ct.build_report(rec, body.get("hits", []), body.get("query", ""))
                    os.startfile(p)
                    return self.send({"name": p.name})
                if u.path == "/api/open-folder":
                    OUT.mkdir(exist_ok=True)
                    os.startfile(OUT)
                    return self.send({"ok": True})
                if u.path == "/api/open-excel":
                    os.startfile(ct.XLSX)
                    return self.send({"ok": True})
                if u.path == "/api/open":
                    target = body.get("url", "")
                    if target.startswith("http") or Path(target).exists():
                        os.startfile(target)
                    return self.send({"ok": True})
                if u.path == "/api/update":
                    start_update("تحديث يدوي")
                    return self.send({"status": "running"})
                if u.path == "/api/docs/download":
                    r = download_docs(int(body["id"]))
                    if any(ok for _n, ok, _x in r["files"]):
                        os.startfile(r["folder"])
                    return self.send(r)
                if u.path == "/api/gazette-account":      # Kuwait Al-Youm subscription (DPAPI-encrypted)
                    import gazette_online as go
                    go.set_account(body.get("user", ""), body.get("password", ""))
                    r = go.test_login() if body.get("password") else {"ok": True, "msg": "اتمسح الاشتراك"}
                    if r.get("ok"):
                        st = gazette_status()
                        st.pop("tried", None)
                        GZ_STATUS.write_text(json.dumps(st, ensure_ascii=False), encoding="utf-8")
                        threading.Thread(target=gazette_watch, args=(lambda m: state["log"].append(m),), daemon=True).start()
                    return self.send(r)
                if u.path == "/api/capt-account":         # the owner types it in the window; kept DPAPI-encrypted
                    import tender_docs as td
                    td.set_account(body.get("user", ""), body.get("password", ""))
                    return self.send(td.test_login() if body.get("password") else {"ok": True, "msg": "اتمسح الحساب"})
                if u.path == "/api/accounts":             # phone app accounts (admin)
                    return self.send(accounts_api(body, lambda m: state["log"].append(m)))
                if u.path == "/api/seen":
                    mark_seen(body.get("tab", ""))
                    return self.send({"ok": True})
                if u.path == "/api/open-file":            # an attachment we saved (Sanjay's reports)
                    import fawaz_watch
                    f = Path(body.get("path", "")).resolve()
                    if f.exists() and str(f).startswith(str(fawaz_watch.WATCH_DIR.resolve())):
                        os.startfile(f)
                    return self.send({"ok": True})
                if u.path == "/api/settings":
                    if "auto_update_minutes" in body:
                        cfg.save_setting("auto_update_minutes", max(0, int(body["auto_update_minutes"])))
                    return self.send({"auto_minutes": auto_minutes()})
                self.send({"error": "not found"}, 404)
            except Exception as e:      # noqa: BLE001
                self.send({"error": str(e)}, 500)

    class Server(ThreadingHTTPServer):
        allow_reuse_address = False

    url = f"http://127.0.0.1:{PORT}/"
    try:
        srv = Server(("127.0.0.1", PORT), H)
    except OSError:
        if open_browser:
            webbrowser.open(url)
        return
    threading.Thread(target=scheduler, daemon=True).start()
    if open_browser:
        threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    srv.serve_forever()


# ----------------------------------------------------------------- CLI
def main(argv):
    if hasattr(sys.stdout, "reconfigure") and sys.stdout:
        sys.stdout.reconfigure(encoding="utf-8")
    if not argv:
        print(__doc__)
        return
    cmd, args, opt = argv[0], [], {}
    it = iter(argv[1:])
    for a in it:
        if a.startswith("--"):
            opt[a[2:]] = next(it, "1")
        else:
            args.append(a)
    q = " ".join(args)
    if cmd == "update":
        only = [c.strip().upper() for c in opt["only"].split(",")] if opt.get("only") else None
        print(json.dumps(run_update(countries=only), ensure_ascii=False, indent=1))
    elif cmd == "search":
        for i, r in enumerate(search(q, limit=int(opt.get("limit", 20)), country=opt.get("country", ""))):
            print(f"{i}) [{r['country']}] [{r['status']}] {r['number']} | {r['org']} | {r['subject'][:100]} | إقفال {r['closing']} | {r['winner'][:40]} {r['value']} | id {r['id']}")
    elif cmd == "report":
        res = search(q, limit=10)
        if res:
            print(build_pdf(res[min(int(opt.get("pick", 0)), len(res) - 1)]["id"]))
    elif cmd == "serve":
        serve(open_browser="nobrowser" not in opt)
    elif cmd == "crypto-test":                    # the accounts' encryption works inside the installed program
        import accounts, tempfile
        res = {}
        try:
            pub, priv = accounts.new_keypair("x")
            d = Path(tempfile.mkdtemp())
            (d / "a.json").write_text("[1,2,3]", encoding="utf-8")
            n = accounts.seal_dir(d, {"users": {"t": {"active": True, "pub": pub, "priv": priv}}})
            res = {"ok": True, "accounts": n, "files": sorted(p.name for p in d.iterdir())}
        except Exception as e:      # noqa: BLE001
            res = {"ok": False, "error": f"{type(e).__name__}: {e}"}
        Path(args[0] if args else "crypto test.json").write_text(json.dumps(res), encoding="utf-8")
    elif cmd == "outlook-test":                   # diagnostics for the installed program (it has no console)
        import fawaz_watch
        m = fawaz_watch.outlook_items()
        Path(args[0] if args else "outlook test.json").write_text(json.dumps(
            {"found": len(m), "error": fawaz_watch.OUTLOOK_ERR[0], "first": m[0]["subject"] if m else "",
             "received": m[0]["received"] if m else ""}, ensure_ascii=False), encoding="utf-8")
    else:
        print(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
