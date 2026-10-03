# -*- coding: utf-8 -*-
"""متابعة فواز و KJAC - everything about our own tenders and the two companies, in one place.

Sources (PC only - this part is company-confidential and never leaves the PC unencrypted):
  - the tendering history workbook from Sanjay (Tenders Report_HR*.xlsx: one sheet per year, every
    tender = a 4-row block: description/client, status/closing, tender id/bidders L1.., bid prices)
  - Sanjay's "tenders on hand" e-mails in Outlook (classic Outlook on this PC, read through COM);
    the attached workbook is saved and parsed (live tenders, closed tenders with results)
Public sources it watches for them: the Gulf directory (status, closing, winner of the same tender
number), the CAPT awards list, the news, plus Google News about Fawaz and KJAC themselves.

Every new thing becomes a row in `feed`; the programs show unread counts from it.
"""
import re, json, sqlite3, datetime, shutil
from pathlib import Path

import fawaz_config as cfg

WATCH_DIR = cfg.ARCHIVE / "متابعة فواز"
ONHAND_DIR = WATCH_DIR / "tenders on hand"
DB = WATCH_DIR / "watch.sqlite"
BASELINE = "2000-01-01T00:00:00"          # what the first run finds is the starting point, not "new"
COMPANY_RX = r"فواز|fawaz|kjac|الكويتي[هة] الياباني[هة]|الكويتية اليابانية|kuwait japan(ese)? air"

SCHEMA = """
CREATE TABLE IF NOT EXISTS tenders(
  key TEXT PRIMARY KEY, src TEXT, year INT, title TEXT, client TEXT, number TEXT, numkey TEXT,
  closing TEXT, status TEXT, bidders TEXT, fawaz_rank INT, l1 TEXT, remarks TEXT, live INT, updated TEXT);
CREATE TABLE IF NOT EXISTS emails(
  id TEXT PRIMARY KEY, received TEXT, sender TEXT, subject TEXT, body TEXT, attachment TEXT);
CREATE TABLE IF NOT EXISTS links(
  tender_key TEXT, kind TEXT, ref TEXT, title TEXT, snap TEXT, url TEXT, first_seen TEXT,
  UNIQUE(tender_key, kind, ref));
CREATE TABLE IF NOT EXISTS feed(
  id INTEGER PRIMARY KEY, at TEXT, kind TEXT, title TEXT, text TEXT, url TEXT, tender_key TEXT, uniq TEXT UNIQUE);
CREATE TABLE IF NOT EXISTS seen(tab TEXT PRIMARY KEY, marker TEXT);
"""


def db():
    WATCH_DIR.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB, timeout=30)
    con.row_factory = sqlite3.Row
    con.executescript(SCHEMA)
    return con


def now():
    return datetime.datetime.now().isoformat(timespec="seconds")


def numkey(tno):
    """digits of a tender number, same rule as the directory: 'و ك م /53/2024/2025' -> '53/2024/2025'"""
    return "/".join(re.findall(r"\d+", str(tno or "")))


def strong_numkey(nk):
    """a number that is safe to match on its own (not '7' or '2026')"""
    parts = nk.split("/") if nk else []
    return len(parts) >= 3 or (len(parts) == 2 and len(nk) >= 7) or (len(parts) == 1 and len(nk) >= 6)


def _date(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    m = re.search(r"(\d{1,2})[./-](\d{1,2})[./-](\d{4})", str(v or ""))
    if m:
        try:
            return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1))).isoformat()
        except ValueError:
            return ""
    return ""


def _s(v):
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def is_us(name):
    return bool(re.search(r"fawaz|فواز|kjac", str(name or ""), re.I))


# ----------------------------------------------------------------- source: history workbook
def history_files():
    files = []
    for d in [WATCH_DIR] + [Path(p) for p in cfg.S.get("watch_history_dirs", [r"D:\Download 2026"])]:
        if d.is_dir():
            files += sorted(d.glob("Tenders Report_HR*.xlsx"), key=lambda f: f.stat().st_mtime)
    return files[-1:] if files else []


def parse_history(path):
    import openpyxl
    out = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        m = re.match(r"Tenders\s*-\s*(\d{4})", ws.title.strip())
        if not m:
            continue
        year = int(m.group(1))
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        head = next((i for i, r in enumerate(rows[:15]) if r and any(_s(x) == "Client" for x in r)), None)
        if head is None:
            continue
        cc = next(j for j, x in enumerate(rows[head]) if _s(x) == "Client")
        starts = [i for i in range(head + 1, len(rows)) if rows[i] and rows[i][0] not in (None, "")
                  and re.fullmatch(r"\d+(\.\d+)?", _s(rows[i][0])) and _s(rows[i][1] if len(rows[i]) > 1 else "")]
        for a, b in zip(starts, starts[1:] + [len(rows)]):
            blk = [r + [None] * (cc + 14 - len(r)) for r in rows[a:min(b, a + 5)]]
            while len(blk) < 4:
                blk.append([None] * (cc + 14))
            title, client = _s(blk[0][1]), _s(blk[0][cc])
            status, closing = _s(blk[1][3]), _date(blk[1][cc])
            number = _s(blk[2][3])
            bidders = []
            for j in range(cc, cc + 13):
                n, p = _s(blk[2][j]), _s(blk[3][j])
                if n and n != "-" and len(n) > 2 and not re.match(r"\d{4}-\d\d-\d\d|[\d.,\s%-]+$", n):
                    bidders.append([n, p])
            rank = next((k + 1 for k, (n, _p) in enumerate(bidders) if is_us(n)), 0)
            if not title:
                continue
            nk = numkey(number) if re.search(r"\d", number) and not re.search(r"mail|email|dt\.?\s|dated", number, re.I) else ""
            key = f"h{year}-{nk or title[:40]}"
            if any(o["key"] == key for o in out):      # several lots under one number
                key += f"-{_s(blk[0][0])}"
            out.append({"key": key, "src": f"سجل {year}", "year": year,
                        "title": title, "client": client, "number": number, "numkey": nk, "closing": closing,
                        "status": status, "bidders": bidders, "fawaz_rank": rank,
                        "l1": bidders[0][0] if bidders else "", "remarks": "", "live": 0})
    wb.close()
    return out


# ----------------------------------------------------------------- source: "tenders on hand" e-mails
def outlook_items(days=150):
    """Sanjay's tenders-on-hand e-mails from the local Outlook (newest first); [] when Outlook is not there"""
    OUTLOOK_ERR[0] = ""
    try:
        import pythoncom
        import win32com.client.dynamic
    except ImportError as e:
        OUTLOOK_ERR[0] = f"pywin32: {e}"
        return []
    pythoncom.CoInitialize()
    try:
        # dynamic dispatch: no makepy / gen_py cache, which is not writable inside the installed program
        ns = win32com.client.dynamic.Dispatch("Outlook.Application").GetNamespace("MAPI")
        items = ns.GetDefaultFolder(6).Items           # Inbox
        since = datetime.datetime.now() - datetime.timedelta(days=days)
        try:
            items = items.Restrict("[ReceivedTime] >= '" + since.strftime("%m/%d/%Y %I:%M %p") + "'")
        except Exception:                               # noqa: BLE001
            pass
        items.Sort("[ReceivedTime]", True)
        who = cfg.S.get("watch_sender", "sanjay").lower()
        subj = cfg.S.get("watch_subject", "tenders on hand").lower()
        out, scanned, first_err = [], 0, ""
        it = items.GetFirst()                           # GetFirst/GetNext: safer than iterating a dynamic collection
        while it is not None:
            cur, it = it, None
            try:
                it = items.GetNext()
            except Exception:                           # noqa: BLE001
                it = None
            scanned += 1
            try:
                s, snd = str(cur.Subject or ""), str(cur.SenderName or "")
                if who not in snd.lower() or subj not in s.lower():
                    continue
                r0 = cur.ReceivedTime                   # Outlook local time, kept as shown in Outlook
                rec = datetime.datetime(r0.year, r0.month, r0.day, r0.hour, r0.minute)
                att = None
                for k in range(1, cur.Attachments.Count + 1):
                    a = cur.Attachments.Item(k)
                    if str(a.FileName).lower().endswith((".xlsx", ".xlsm")):
                        ONHAND_DIR.mkdir(parents=True, exist_ok=True)
                        safe = re.sub(r'[\\/:*?<>|]', "_", str(a.FileName))
                        dst = ONHAND_DIR / f"{rec:%Y-%m-%d %H%M} - {safe}"
                        if not dst.exists():
                            a.SaveAsFile(str(dst))
                        att = dst
                        break
                out.append({"id": str(cur.EntryID), "received": f"{rec:%Y-%m-%dT%H:%M}", "sender": snd, "subject": s,
                            "body": re.sub(r"\n{3,}", "\n\n", str(cur.Body or "").replace("\r", ""))[:2500],
                            "attachment": str(att) if att else ""})
            except Exception as e:                      # noqa: BLE001 - skip an odd item, remember why
                first_err = first_err or f"{type(e).__name__}: {e}"[:200]
        if not out:
            OUTLOOK_ERR[0] = f"اتفحص {scanned} إيميل" + (f" - أول خطأ: {first_err}" if first_err else "")
        return out
    except Exception as e:                              # noqa: BLE001 - Outlook closed / not configured
        OUTLOOK_ERR[0] = str(e)[:200]
        return []
    finally:
        pythoncom.CoUninitialize()


OUTLOOK_ERR = [""]


def parse_onhand(path):
    import openpyxl
    out = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        name = ws.title.strip().lower()
        live = name.startswith("new tenders on hand") or name.startswith("enquiries")
        closed = name.startswith("closed")
        if not (live or closed):
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        hi = next((i for i, r in enumerate(rows[:6]) if r and any(_s(x).lower() == "client" for x in r)), None)
        if hi is None:
            continue
        h = [_s(x).lower() for x in rows[hi]]
        col = lambda *names: next((j for j, x in enumerate(h) if any(x.startswith(n) for n in names)), None)
        c_client, c_title = col("client"), col("new tenders on hand", "tender title", "enquiry subject", "tender subject")
        c_no, c_close, c_rem = col("tender #", "enq ref", "tender no"), col("closing date", "closed date"), col("remarks")
        c_res = (c_rem + 1) if c_rem is not None else None
        i = hi + 1
        while i < len(rows):
            r = rows[i] + [None] * 40
            if not _s(r[c_title] if c_title is not None else "") or not re.fullmatch(r"\d+", _s(r[0])):
                i += 1
                continue
            bidders = []
            if closed and c_res is not None:
                names = [_s(x) for x in r[c_res:c_res + 18]]
                nxt = (rows[i + 1] + [None] * 40) if i + 1 < len(rows) else [None] * 40
                prices = [_s(x) for x in nxt[c_res:c_res + 18]] if not re.fullmatch(r"\d+", _s(nxt[0])) else [""] * 18
                bidders = [[n, p] for n, p in zip(names, prices) if n]
            number = _s(r[c_no]) if c_no is not None else ""
            nk = numkey(number) if re.search(r"\d", number) and not re.search(r"mail|ltr|dated", number, re.I) else ""
            title = _s(r[c_title])
            out.append({"key": f"o-{nk or (_s(r[c_client]) + title)[:50]}", "src": ws.title.strip(), "year": 0,
                        "title": title, "client": _s(r[c_client]) if c_client is not None else "", "number": number,
                        "numkey": nk, "closing": _date(r[c_close]) if c_close is not None else "",
                        "status": "مطروحة - بنجهزها" if live else ("نتيجة" if bidders else "مقفولة"),
                        "bidders": bidders, "fawaz_rank": next((k + 1 for k, (n, _p) in enumerate(bidders) if is_us(n)), 0),
                        "l1": bidders[0][0] if bidders else "", "remarks": _s(r[c_rem]) if c_rem is not None else "",
                        "live": int(live)})
            i += 1
    wb.close()
    return out


# ----------------------------------------------------------------- feed helpers
def add_feed(con, kind, title, text="", url="", tender_key="", uniq=None, at=None):
    cur = con.execute("INSERT OR IGNORE INTO feed(at,kind,title,text,url,tender_key,uniq) VALUES(?,?,?,?,?,?,?)",
                      (at or now(), kind, title, text, url, tender_key, uniq or f"{kind}|{tender_key}|{title}|{text}"[:400]))
    return cur.rowcount


def upsert_tender(con, t, first):
    old = con.execute("SELECT * FROM tenders WHERE key=?", (t["key"],)).fetchone()
    vals = dict(t, bidders=json.dumps(t["bidders"], ensure_ascii=False), updated=now())
    if old is None:
        con.execute("INSERT INTO tenders(" + ",".join(vals) + ") VALUES(" + ",".join("?" * len(vals)) + ")", tuple(vals.values()))
        if t["live"]:
            add_feed(con, "مناقصة بنجهزها", t["title"], f"{t['client']} · {t['number']} · إقفال {t['closing'] or '-'}",
                     tender_key=t["key"], at=BASELINE if first else None)
        return
    changes = []
    if t["closing"] and t["closing"] != old["closing"]:
        changes.append(f"الإقفال اتغيّر من {old['closing'] or '-'} لـ {t['closing']}")
    if t["status"] and t["status"] != old["status"]:
        changes.append(f"الحالة: {t['status']}")
    if t["bidders"] and vals["bidders"] != old["bidders"]:
        top = " · ".join(f"L{k + 1} {n} {p}" for k, (n, p) in enumerate(t["bidders"][:4]))
        changes.append(f"الأسعار: {top}" + (f" · ترتيبنا L{t['fawaz_rank']}" if t["fawaz_rank"] else ""))
    if t["remarks"] and t["remarks"] != old["remarks"]:
        changes.append(f"ملاحظات: {t['remarks']}")
    for k in ("closing", "status", "remarks", "l1", "bidders", "fawaz_rank"):   # never erase what we knew
        if not t[k] and old[k]:
            vals[k] = old[k]
    con.execute("UPDATE tenders SET " + ",".join(f"{k}=?" for k in vals) + " WHERE key=?", (*vals.values(), t["key"]))
    for c in changes:
        add_feed(con, "تحديث من تقرير Sanjay", t["title"], c, tender_key=t["key"])


# ----------------------------------------------------------------- matching against the public data
CLIENTS_AR = {"MEW": "وزارة الكهرباء", "MOH": "وزارة الصحة", "KNPC": "البترول الوطنية", "KOC": "نفط الكويت",
              "MOI": "وزارة الداخلية", "PAHW": "الرعاية السكنية", "KU": "جامعة الكويت", "MPW": "وزارة الأشغال",
              "KM": "بلدية الكويت", "PAAET": "للتعليم التطبيقي", "MOSAL": "الشؤون الاجتماعية", "KNA": "مجلس الأمة",
              "MOE": "وزارة التربية", "KPC": "مؤسسة البترول الكويتية", "KIPIC": "للصناعات البترولية المتكاملة",
              "PADA": "لشؤون ذوي الإعاقة", "KAC": "الخطوط الجوية الكويتية", "DGCA": "الطيران المدني",
              "MOD": "وزارة الدفاع", "KISR": "معهد الكويت للأبحاث"}


def _letters(tno, norm):
    return re.sub(r"[^a-zء-ي]", "", norm(str(tno or "")))


def same_tender(t, number, org, norm):
    """same number is not enough - every ministry restarts at 1/2025/2026: the letters of the number
    (و ك م = MEW) must agree, or the client has to be the same body"""
    nk = t["numkey"]
    if len(nk.split("/")) == 1 and len(nk) >= 6:          # RFP-2143558 style numbers are unique
        return True
    la, lb = _letters(t["number"], norm), _letters(number, norm)
    if la and lb:
        return la == lb or (min(len(la), len(lb)) >= 2 and (la in lb or lb in la))
    ar = CLIENTS_AR.get((t["client"] or "").upper().strip())
    return bool(ar and norm(ar) in norm(org or ""))


def match_public(con, gd, ct, first, progress=print):
    """link our tenders to directory items / CAPT awards with the same tender number; report what changed"""
    tenders = [dict(r) for r in con.execute("SELECT * FROM tenders WHERE numkey<>''")]
    by_nk = {}
    for t in tenders:
        if strong_numkey(t["numkey"]):
            by_nk.setdefault(t["numkey"], []).append(t)
    n = 0
    dcon = gd.db()
    for it in dcon.execute("SELECT id, numkey, number, org, subject, status, closing, winner, value, official FROM items WHERE numkey<>''"):
        for t in by_nk.get(it["numkey"], []):
            if not same_tender(t, it["number"], it["org"], gd.norm):
                continue
            snap = json.dumps([it["status"], it["closing"], it["winner"], it["value"]], ensure_ascii=False)
            old = con.execute("SELECT snap FROM links WHERE tender_key=? AND kind='item' AND ref=?", (t["key"], str(it["id"]))).fetchone()
            if old is None:
                con.execute("INSERT INTO links VALUES(?,?,?,?,?,?,?)", (t["key"], "item", str(it["id"]), it["subject"], snap,
                                                                     it["official"] or "", now()))
                add_feed(con, "لقيناها في الدليل", t["title"], f"{it['org']} · {it['status']}" +
                         (f" · الفائز {it['winner']} {it['value']}" if it["winner"] else ""), it["official"] or "",
                         t["key"], at=BASELINE if first else None)
                n += 1
            elif old["snap"] != snap:
                o = json.loads(old["snap"])
                diff = [f"{lbl}: {a or '-'} ← {b}" for lbl, a, b in zip(("الحالة", "الإقفال", "الفائز", "القيمة"), o,
                                                                        [it["status"], it["closing"], it["winner"], it["value"]]) if b and a != b]
                con.execute("UPDATE links SET snap=? WHERE tender_key=? AND kind='item' AND ref=?", (snap, t["key"], str(it["id"])))
                if diff:
                    add_feed(con, "اتغيّرت في الدليل", t["title"], " · ".join(diff), it["official"] or "", t["key"])
                    n += 1
    # gazette / CAPT-decision events on those items
    for lk in con.execute("SELECT tender_key, ref FROM links WHERE kind='item'").fetchall():
        for e in dcon.execute("SELECT id, date, source, etype, text FROM events WHERE item_id=?", (int(lk["ref"]),)):
            if con.execute("SELECT 1 FROM links WHERE tender_key=? AND kind='event' AND ref=?", (lk["tender_key"], str(e["id"]))).fetchone():
                continue
            con.execute("INSERT INTO links VALUES(?,?,?,?,?,?,?)", (lk["tender_key"], "event", str(e["id"]), e["etype"], "", "", now()))
            t = con.execute("SELECT title FROM tenders WHERE key=?", (lk["tender_key"],)).fetchone()
            add_feed(con, e["source"] or "الدليل", t["title"] if t else "", f"{e['date']} · {e['etype']} · {e['text'][:300]}",
                     tender_key=lk["tender_key"], at=BASELINE if first else None)
    dcon.close()
    # CAPT awards with the same number
    try:
        for a in ct.read_excel_rows():
            nk = numkey(a["tno"])
            for t in by_nk.get(nk, []):
                if not same_tender(t, a["tno"], a["org"], gd.norm):
                    continue
                ref = f"{a['date']}|{a['winner']}|{a['total']}"
                if con.execute("SELECT 1 FROM links WHERE tender_key=? AND kind='award' AND ref=?", (t["key"], ref)).fetchone():
                    continue
                con.execute("INSERT INTO links VALUES(?,?,?,?,?,?,?)", (t["key"], "award", ref, a["subject"], "", "", now()))
                url = f"{ct.SITE}/ar/tenders/winning-bids/?meeting_date_from={a['date']}&meeting_date_to={a['date']}&form=date"
                add_feed(con, "ترسية من الجهاز المركزي", t["title"], f"{a['date']} · {a['winner'] or 'بنود'} · {a['total']}",
                         url, t["key"], at=BASELINE if first else None)
                n += 1
    except Exception as e:                              # noqa: BLE001
        progress(f"متابعة فواز - الترسيات: {e}")
    return n


def match_news(con, gd, first):
    """news that names one of our live / recent tenders by number, or the client + our kind of work"""
    recent = (datetime.date.today() - datetime.timedelta(days=400)).isoformat()
    tenders = [dict(r) for r in con.execute("SELECT * FROM tenders WHERE live=1 OR closing>=?", (recent,))]
    dcon = gd.db()
    news = [dict(r) for r in dcon.execute("SELECT id, date, source, title, summary, url FROM news WHERE date>=? AND country='KW'", (recent,))]
    dcon.close()
    live_by_client = {}
    for t in tenders:
        ar = CLIENTS_AR.get((t["client"] or "").upper().strip())
        if t["live"] and ar:
            live_by_client.setdefault(gd.norm(ar), []).append(t)
    n = 0
    for nw in news:
        text = gd.norm(nw["title"] + " " + (nw["summary"] or ""))
        digits = set(re.findall(r"\d+(?:/\d+)+", text.replace(" ", "")))
        for t in tenders:                               # the tender number itself is in the news
            if t["numkey"] and strong_numkey(t["numkey"]) and t["numkey"] in digits and \
                    con.execute("INSERT OR IGNORE INTO links VALUES(?,?,?,?,?,?,?)",
                                (t["key"], "news", str(nw["id"]), nw["title"], "", nw["url"], now())).rowcount:
                add_feed(con, "خبر يخص مناقصة لينا", t["title"], f"{nw['date']} · {nw['source']} · {nw['title']}",
                         nw["url"], t["key"], at=BASELINE if first else None)
                n += 1
        # maintenance / HVAC news about a client we are bidding with right now - once per news item
        if re.search(r"صيان|تشغيل|تكييف|تبريد|مناقص|ترسي", text):
            for ar, lst in live_by_client.items():
                if ar in text:
                    n += add_feed(con, "خبر عن عميل عندنا معاه مناقصة شغالة", nw["title"],
                                  f"{nw['date']} · {nw['source']} · عندنا {len(lst)} مناقصة شغالة مع الجهة ده: " +
                                  " ؛ ".join(x["title"][:60] for x in lst[:4]), nw["url"],
                                  uniq=f"clientnews|{nw['url']}", at=BASELINE if first else None)
                    break
    return n


def company_news_feed(con, gd, first):
    """news about Fawaz / KJAC themselves (the directory's news step), plus any news or CAPT decision /
    gazette notice in the directory that names one of the two companies"""
    dcon = gd.db()
    rows = dcon.execute("SELECT id, date, source, title, url FROM news WHERE query LIKE 'شركة:%' ORDER BY date").fetchall()
    named = [r for r in dcon.execute("SELECT id, date, source, title, summary, url FROM news WHERE query NOT LIKE 'شركة:%'")
             if re.search(COMPANY_RX, (r["title"] or "") + " " + (r["summary"] or ""), re.I)]
    events = [r for r in dcon.execute("SELECT e.id, e.date, e.source, e.etype, e.text, e.url, i.subject, i.official FROM events e "
                                      "JOIN items i ON i.id=e.item_id WHERE e.source NOT LIKE 'تقرير مشاريع%'")
              if re.search(COMPANY_RX, e_text(r), re.I)]
    dcon.close()
    n = 0
    for r in rows:
        n += add_feed(con, "خبر عن فواز / KJAC", r["title"], f"{r['date']} · {r['source']}", r["url"],
                      uniq=f"conews|{r['url']}", at=BASELINE if first else None)
    for r in named:
        n += add_feed(con, "خبر بيذكر فواز / KJAC", r["title"], f"{r['date']} · {r['source']}", r["url"],
                      uniq=f"conews|{r['url']}", at=BASELINE if first else None)
    for r in events:
        n += add_feed(con, "قرار / إعلان بيذكر فواز / KJAC", r["subject"], f"{r['date']} · {r['source']} · {r['etype']} · {r['text'][:300]}",
                      r["official"] if (r["official"] or "").startswith("http") else "",
                      uniq=f"coevent|{r['id']}", at=BASELINE if first else None)
    return n


def e_text(r):
    return f"{r['etype'] or ''} {r['text'] or ''}"


# ----------------------------------------------------------------- run
def run(gd, ct, progress=print):
    """called after every directory update on the PC"""
    con = db()
    first = con.execute("SELECT count(*) FROM tenders").fetchone()[0] == 0
    summary = {}
    for f in history_files():
        rows = parse_history(f)
        for t in rows:
            upsert_tender(con, t, first)
        summary["سجل المناقصات"] = len(rows)
        if f.parent != WATCH_DIR:                       # keep our own copy next to the data
            try:
                WATCH_DIR.mkdir(parents=True, exist_ok=True)
                shutil.copy2(f, WATCH_DIR / f.name)
            except OSError:
                pass
    mails = outlook_items()
    new_mail = 0
    for m in reversed(mails):
        if con.execute("SELECT 1 FROM emails WHERE id=?", (m["id"],)).fetchone():
            continue
        con.execute("INSERT INTO emails VALUES(?,?,?,?,?,?)", tuple(m[k] for k in ("id", "received", "sender", "subject", "body", "attachment")))
        add_feed(con, "إيميل Sanjay", m["subject"], f"{m['received'].replace('T', ' ')} · {Path(m['attachment']).name if m['attachment'] else 'من غير مرفق'}",
                 uniq=f"mail|{m['id']}", at=BASELINE if first else None)
        new_mail += 1
    summary["إيميلات Sanjay"] = f"{len(mails)} ({new_mail} جديد)" if mails else \
        f"مفيش من Outlook ({OUTLOOK_ERR[0] or 'ملقيتش إيميلات tenders on hand آخر 150 يوم'})"
    latest = next((m["attachment"] for m in mails if m["attachment"]), "")
    if not latest:
        old = sorted(ONHAND_DIR.glob("*.xls*")) if ONHAND_DIR.exists() else []
        latest = str(old[-1]) if old else ""
    if latest:
        rows = parse_onhand(latest)
        live_keys = {t["key"] for t in rows if t["live"]}
        con.execute("UPDATE tenders SET live=0 WHERE src NOT LIKE 'سجل%' AND key NOT IN (" +
                    ",".join("?" * len(live_keys)) + ")", tuple(live_keys))
        for t in rows:
            upsert_tender(con, t, first)
        summary["آخر تقرير tenders on hand"] = f"{Path(latest).name}: {len(live_keys)} شغالة"
    con.commit()
    summary["ربط بالدليل والترسيات"] = match_public(con, gd, ct, first, progress)
    summary["أخبار مناقصاتنا"] = match_news(con, gd, first)
    summary["أخبار الشركة"] = company_news_feed(con, gd, first)
    con.commit()
    con.close()
    progress(f"متابعة فواز: {summary}")
    return summary


# ----------------------------------------------------------------- read side (window + phone)
def data(limit_feed=400):
    if not DB.exists():
        return {"feed": [], "tenders": [], "emails": [], "seen": {}}
    con = db()
    feed = [dict(r) for r in con.execute("SELECT id, at, kind, title, text, url, tender_key FROM feed ORDER BY id DESC LIMIT ?", (limit_feed,))]
    tenders = []
    for r in con.execute("SELECT * FROM tenders ORDER BY live DESC, closing DESC"):
        t = dict(r)
        t["bidders"] = json.loads(t["bidders"] or "[]")
        t["links"] = [dict(x) for x in con.execute("SELECT kind, ref, title, url, first_seen FROM links WHERE tender_key=? AND kind<>'event' ORDER BY first_seen DESC", (t["key"],))]
        tenders.append(t)
    emails = [dict(r) for r in con.execute("SELECT id, received, sender, subject, body, attachment FROM emails ORDER BY received DESC LIMIT 40")]
    seen = dict(con.execute("SELECT tab, marker FROM seen").fetchall())
    con.close()
    return {"feed": feed, "tenders": tenders, "emails": emails, "seen": seen}


def get_seen():
    if not DB.exists():
        return {}
    con = db()
    s = dict(con.execute("SELECT tab, marker FROM seen").fetchall())
    con.close()
    return s


def set_seen(tab, marker):
    con = db()
    con.execute("INSERT OR REPLACE INTO seen VALUES(?,?)", (tab, str(marker)))
    con.commit()
    con.close()


def unread_feed(marker):
    if not DB.exists():
        return 0
    con = db()
    n = con.execute("SELECT count(*) FROM feed WHERE id>? AND at>?", (int(marker or 0), BASELINE)).fetchone()[0]
    con.close()
    return n
