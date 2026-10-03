# -*- coding: utf-8 -*-
"""أداة ترسيات CAPT - FAWAZ

  python "capt tool.py" update            تحديث الإكسيل من الموقع + تحميل المحاضر الجديدة + فهرستها
  python "capt tool.py" index             فهرسة ملفات المحاضر (الجديد أو المتغير بس)
  python "capt tool.py" search "<نص>"     بحث في الترسيات (الإكسيل)
  python "capt tool.py" minutes "<نص>"    بحث مباشر في المحاضر
  python "capt tool.py" report "<نص>" [--pick N] [--row R]
                                          PDF واحد: بيانات الترسية + صفحات المحاضر المتعلقة
  python "capt tool.py" serve             واجهة البرنامج (بتفتح في المتصفح)

The site (capt.gov.kw) is behind Cloudflare but answers plain requests that carry a
browser User-Agent. Minutes PDFs text uses Arabic presentation forms -> NFKC, and the
numbers inside them are often bidi-scrambled (53/2024/2025 shows as 2025/2024/53), so
tender numbers are matched in both directions / by FTS NEAR.
"""
import sys, os, re, json, sqlite3, unicodedata, datetime, shutil, time, math, html
import urllib.request, collections, threading
from pathlib import Path

import fawaz_config as cfg
import updater
BASE = cfg.CAPT_DIR                                    # أداة ترسيات CAPT (index, reports, backups)
ARCHIVE = cfg.ARCHIVE                                  # التسعير والعطاءات
XLSX = cfg.AWARDS_XLSX
LAST_RUN = BASE / "last update.txt"
MOM = ARCHIVE / "CAPT MOM"
NEW_MOM = MOM / "محاضر جديدة من الموقع"
DB = BASE / "capt index.sqlite"
OUT = BASE / "تقارير المناقصات"
BACKUP = BASE / "نسخ احتياطية"
LOG = BASE / "سجل التحديث.txt"
MEETINGS = BASE / "meetings.json"
EXCEL_CACHE = BASE / "excel cache.json"
SITE = "https://capt.gov.kw"
PORT = 8765
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/140.0 Safari/537.36",
      "Accept-Language": "ar,en;q=0.9", "Accept": "text/html,application/xhtml+xml,*/*"}
FONTS = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"

# ----------------------------------------------------------------- text helpers
_TR = str.maketrans({"أ": "ا", "إ": "ا", "آ": "ا", "ٱ": "ا", "ى": "ي", "ة": "ه", "ؤ": "و",
                     "ئ": "ي", "ک": "ك", "ی": "ي", "٠": "0", "١": "1", "٢": "2", "٣": "3",
                     "٤": "4", "٥": "5", "٦": "6", "٧": "7", "٨": "8", "٩": "9",
                     "۰": "0", "۱": "1", "۲": "2", "۳": "3", "۴": "4", "۵": "5",
                     "۶": "6", "۷": "7", "۸": "8", "۹": "9"})
_DIAC = re.compile(r"[\u064B-\u065F\u0670\u06D6-\u06ED\u0640]")


def norm(s):
    s = unicodedata.normalize("NFKC", str(s or ""))
    s = _DIAC.sub("", s).translate(_TR).lower()
    return re.sub(r"\s+", " ", s).strip()


STOP = set(norm(w) for w in """في من على الى إلى عن مع او أو و ثم بشأن بشان لمدة لصالح لزوم
التابعة التابعه التابع الخاصة الخاصه وذلك ذلك هذه هذا تلك التي الذي كافة كافه جميع بعض عدد
لسنة لسنه سنة سنه رقم طرح مناقصة مناقصه عامة عامه الكويت دولة دوله بدولة بدوله مختلفة متفرقة
ل ب ك the of and for to in at""".split())


def stem(w):
    for p in ("وال", "بال", "كال", "فال", "لل", "ال"):
        if w.startswith(p) and len(w) - len(p) >= 2:
            return w[len(p):]
    return w


def words(s):
    return re.findall(r"[a-z0-9]+|[\u0621-\u064a]+", norm(s))


def stems(s):
    return [stem(w) for w in words(s) if w not in STOP and len(w) > 1]


def num_parts(tno):
    """digit groups of a tender number: 'و ك م /53/2024/2025' -> ['53','2024','2025']"""
    return re.findall(r"\d+", norm(tno))


def is_year(n):
    return len(n) == 4 and n[:2] in ("19", "20")


def key_nums(tno):
    return [n for n in num_parts(tno) if not is_year(n)]


AR_MONTHS = {"يناير": 1, "فبراير": 2, "مارس": 3, "ابريل": 4, "مايو": 5, "يونيو": 6,
             "يوليو": 7, "اغسطس": 8, "سبتمبر": 9, "اكتوبر": 10, "نوفمبر": 11, "ديسمبر": 12}


def parse_ar_date(s):
    """'سبتمبر 23, 2026' or '23 سبتمبر،2026' -> date"""
    s = norm(s)
    m = re.search(r"([\u0621-\u064a]+)\s*(\d{1,2})\s*[,،]?\s*(\d{4})", s) or \
        re.search(r"(\d{1,2})\s*([\u0621-\u064a]+)\s*[,،]?\s*(\d{4})", s)
    if not m:
        return None
    a, b, y = m.groups()
    mon, day = (a, b) if not a.isdigit() else (b, a)
    if mon not in AR_MONTHS:
        return None
    return datetime.date(int(y), AR_MONTHS[mon], int(day))


def to_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    if isinstance(v, str):
        m = re.match(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})", v)
        if m:
            try:
                return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
            except ValueError:
                return None
    return None


def dmy(d):
    return d.strftime("%d/%m/%Y") if d else ""


def log(msg):
    line = f"[{datetime.datetime.now():%Y-%m-%d %H:%M}] {msg}"
    print(line, flush=True)
    try:
        with open(LOG, "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except OSError:
        pass


def http_get(url, binary=False, tries=4):
    last = None
    for k in range(tries):
        try:
            req = urllib.request.Request(url, headers=UA)
            with urllib.request.urlopen(req, timeout=90) as r:
                data = r.read()
            return data if binary else data.decode("utf-8", "replace")
        except Exception as e:          # noqa: BLE001 - retry any network error
            last = e
            time.sleep(3 * (k + 1))
    raise RuntimeError(f"فشل تحميل {url}: {last}")


def strip_tags(s):
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", s))).strip()


# ----------------------------------------------------------------- site: awards
def parse_awards_page(h):
    i = h.find('class="custom-table"')
    if i < 0:
        return [], 1
    j = h.find("el-pagination", i)
    body = h[i:j if j > 0 else len(h)]
    rows = []
    for chunk in body.split('class="table-row tbody"')[1:]:
        cells = [strip_tags(c) for c in re.findall(r'<div class="table-cell">(.*?)</div>', chunk, re.S)]
        pop = re.search(r'data-popup-url="([^"]+)"', chunk)
        if len(cells) < 5:
            continue
        r = {"item": cells[0], "date": parse_ar_date(cells[1]), "tno": cells[2], "subject": cells[3],
             "org": cells[4], "winner": "", "total": "", "notes": "", "popup": pop.group(1) if pop else ""}
        if pop:
            r["notes"] = "بنود"
        else:
            r["winner"] = cells[5] if len(cells) > 5 else ""
            r["total"] = cells[6] if len(cells) > 6 else ""
            r["notes"] = cells[7] if len(cells) > 7 else ""
        rows.append(r)
    pages = [int(x) for x in re.findall(r"[?&]page=(\d+)", h)]
    return rows, max(pages) if pages else 1


def fetch_popup(url):
    h = http_get(SITE + url)
    items = []
    for tr in re.findall(r"<tr>(.*?)</tr>", h.split("tenderpoptable2", 1)[-1], re.S):
        tds = [strip_tags(x) for x in re.findall(r"<td[^>]*>(.*?)</td>", tr, re.S)]
        if len(tds) >= 3:
            items.append(tds[:3])
    return items


def fmt_kd(v):
    return f"{v:,.3f} د.ك"


def fill_items(r):
    """rows that only have a 'بنود' button: winners + total come from the popup"""
    try:
        items = fetch_popup(r["popup"])
    except RuntimeError:
        return
    if not items:
        return
    winners, total, parts = [], 0.0, []
    for w, no, price in items:
        if w and w not in winners:
            winners.append(w)
        m = re.search(r"-?[\d,]+(?:\.\d+)?", price.replace(" ", ""))
        if m:
            try:
                total += float(m.group(0).replace(",", ""))
            except ValueError:
                pass
        parts.append(f"{w} (بند {no}: {price})")
    r["winner"] = " + ".join(winners)
    r["total"] = fmt_kd(total) if total else ""
    shown = parts[:12]
    r["notes"] = f"بنود ({len(items)}): " + " ؛ ".join(shown) + (" ..." if len(parts) > 12 else "")


def row_key(date, tno, winner, total, notes):
    if "بنود" in str(notes or "") and not str(winner or "").strip():
        return (date, re.sub(r"\s", "", norm(tno)), "بنود")
    return (date, re.sub(r"\s", "", norm(tno)), re.sub(r"\s", "", norm(winner)),
            re.sub(r"\s", "", norm(total)))


def fetch_day(day):
    q = f"meeting_date_from={day:%Y-%m-%d}&meeting_date_to={day:%Y-%m-%d}&form=date"
    url = f"{SITE}/ar/tenders/winning-bids/?{q}"
    rows, last = parse_awards_page(http_get(url))
    if last <= 1:
        return rows
    # the site's order inside one day is not stable between requests, so page through
    # several times until every row has been seen
    seen = {}
    for r in rows:
        seen[(r["item"], r["tno"], r["winner"], r["total"], r["popup"])] = r
    expected = None
    for _ in range(6):
        for p in range(1, last + 1):
            rr, _l = parse_awards_page(http_get(f"{url}&page={p}"))
            if p == last:
                expected = 10 * (last - 1) + len(rr)
            for r in rr:
                seen[(r["item"], r["tno"], r["winner"], r["total"], r["popup"])] = r
        if expected and len(seen) >= expected:
            break
    return list(seen.values())


def fetch_awards_since(d0, d1=None):
    d1 = d1 or datetime.date.today()
    out, day = [], d1
    while day >= d0:
        rows = fetch_day(day)
        rows.sort(key=lambda r: int(r["item"]) if r["item"].isdigit() else 999)
        out.extend(rows)
        day -= datetime.timedelta(days=1)
    for r in out:
        if r["popup"]:
            fill_items(r)
    return out


# ----------------------------------------------------------------- excel
def read_excel_rows():
    """all award rows of sheet All (cached by file mtime)"""
    st = XLSX.stat()
    sig = f"v2-{st.st_mtime_ns}-{st.st_size}"
    if EXCEL_CACHE.exists():
        try:
            c = json.loads(EXCEL_CACHE.read_text(encoding="utf-8"))
            if c.get("sig") == sig:
                return c["rows"]
        except (ValueError, KeyError):
            pass
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    ws = wb["All"]
    rows = []
    for i, r in enumerate(ws.iter_rows(min_row=3, max_col=9, values_only=True), start=3):
        if not r or len(r) < 9 or not r[3]:
            continue
        d = to_date(r[6])
        winner, notes = clean_winner(str(r[5] or "").strip(), str(r[8] or "").strip())
        rows.append({"row": i, "sr": r[1], "tno": str(r[2] or "").strip(), "subject": str(r[3]).strip(),
                     "org": str(r[4] or "").strip(), "winner": winner,
                     "date": d.isoformat() if d else "", "total": fmt_total(r[7]), "notes": notes})
    wb.close()
    try:
        EXCEL_CACHE.write_text(json.dumps({"sig": sig, "rows": rows}, ensure_ascii=False), encoding="utf-8")
    except OSError:
        pass
    return rows


def clean_winner(w, notes):
    """old item-type rows hold the whole site popup text in Winner -> company names only"""
    if "Winning Bidders" not in w:
        return w, notes
    names = []
    for line in w.split("Winning Bidders", 1)[1].splitlines()[1:]:
        name = re.sub(r"\s+\d+\s+(?:[\d.,]+|-+)\s*KD\s*$", "", line.strip()).strip()
        if name and name not in names:
            names.append(name)
    items = [ln for ln in w.splitlines() if ln.strip().endswith("KD")]
    extra = f"بنود ({len(items)} بند)"
    return " + ".join(names), (notes if notes and notes != "ITEMS" else extra)


def fmt_total(v):
    if v is None:
        return ""
    if isinstance(v, float) and abs(v) < 1:          # stored discount percentages
        return f"{v * 100:.2f}%".replace(".00%", "%")
    if isinstance(v, (int, float)):
        return fmt_kd(v)
    return str(v).strip()


def update_excel(progress=print):
    import openpyxl
    from copy import copy
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["All"]
    existing, last = set(), None
    for r in range(3, ws.max_row + 1):
        if ws.cell(r, 4).value is None:
            continue
        d = to_date(ws.cell(r, 7).value)
        if d and (last is None or d > last):
            last = d
    if last is None:
        raise RuntimeError("مش لاقي آخر تاريخ ترسية في الشيت All")
    for r in range(3, ws.max_row + 1):
        d = to_date(ws.cell(r, 7).value)
        if d and d >= last:
            existing.add(row_key(d, *(ws.cell(r, c).value for c in (3, 6, 8, 9))))
    progress(f"آخر ترسية في الملف: {dmy(last)} - بجيب من الموقع لحد النهارده ...")
    fetched = fetch_awards_since(last)
    new = [r for r in fetched
           if row_key(r["date"], r["tno"], r["winner"], r["total"], r["notes"]) not in existing
           and not (r["notes"].startswith("بنود") and
                    row_key(r["date"], r["tno"], "", "", "بنود") in existing)]
    n = len(new)
    if not n:                       # nothing new: leave the file (and its backups) alone
        return new, last
    # the pasted top rows keep dates as dd/mm/yyyy text -> make them real dates, same look
    for r in range(3, ws.max_row + 1):
        v = ws.cell(r, 7).value
        if isinstance(v, str) and to_date(v):
            ws.cell(r, 7).value = datetime.datetime.combine(to_date(v), datetime.time())
            ws.cell(r, 7).number_format = "dd/mm/yyyy"
    if n:
        # Sr No block at the top counts 1,2,3... -> shift it
        blk_end, prev = 3, None
        for r in range(3, ws.max_row + 1):
            s = ws.cell(r, 2).value
            if not isinstance(s, int) or (prev is not None and s != prev + 1):
                break
            prev, blk_end = s, r
        ws.insert_rows(3, n)
        tmpl = [ws.cell(3 + n, c) for c in range(1, 10)]
        for r in range(3 + n, blk_end + n + 1):
            if isinstance(ws.cell(r, 2).value, int):
                ws.cell(r, 2).value += n
        for k, rec in enumerate(new):
            r = 3 + k
            vals = [None, k + 1, rec["tno"], rec["subject"], rec["org"], rec["winner"] or None,
                    datetime.datetime.combine(rec["date"], datetime.time()), rec["total"] or None,
                    rec["notes"] or None]
            for c in range(1, 10):
                cell = ws.cell(r, c)
                cell.value = vals[c - 1]
                t = tmpl[c - 1]
                if t.has_style:
                    cell.font, cell.border, cell.fill = copy(t.font), copy(t.border), copy(t.fill)
                    cell.alignment, cell.protection = copy(t.alignment), copy(t.protection)
                cell.number_format = "dd/mm/yyyy" if c == 7 else t.number_format
    ws.auto_filter.ref = f"B2:I{ws.max_row}"
    BACKUP.mkdir(exist_ok=True)
    bk = BACKUP / f"CAPT Winning Bids from 2006 - قبل تحديث {datetime.date.today():%Y-%m-%d}.xlsx"
    shutil.copy2(XLSX, bk)
    for old in sorted(BACKUP.glob("CAPT Winning Bids*.xlsx"))[:-8]:
        old.unlink()
    tmp = XLSX.with_name("~tmp capt update.xlsx")
    wb.save(tmp)
    try:
        os.replace(tmp, XLSX)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise RuntimeError("ملف الإكسيل مفتوح - اقفله وشغّل التحديث تاني")
    return new, last


# ----------------------------------------------------------------- site: minutes
def fetch_meetings():
    h = http_get(SITE + "/ar/")
    out = {}
    for m in re.finditer(r'<tr class="tr-body[^"]*"[^>]*?data-number="(\d+)/(\d{4})"\s*data-date="([\d-]+)"'
                         r'.*?</tr>', h, re.S):
        no, yr, d = m.groups()
        tr = m.group(0)
        url = re.search(r"location\.href='([^']+\.pdf)'", tr)
        spans = re.findall(r"<span>([^<]+)</span>", tr)
        pub = parse_ar_date(spans[1]) if len(spans) > 1 else None
        out[f"{yr}/{int(no)}"] = {"date": d, "publish": pub.isoformat() if pub else "",
                                  "url": url.group(1) if url else ""}
    if out:
        old = {}
        if MEETINGS.exists():
            old = json.loads(MEETINGS.read_text(encoding="utf-8"))
        old.update(out)
        MEETINGS.write_text(json.dumps(old, ensure_ascii=False, indent=0), encoding="utf-8")
    return out


def load_meetings():
    if MEETINGS.exists():
        return json.loads(MEETINGS.read_text(encoding="utf-8"))
    return {}


def download_new_minutes(progress=print):
    meets = fetch_meetings()
    con = db()
    # a meeting counts as present only if it has real text (the scanned part of 2023 does not)
    have = set(f"{y}/{m}" for y, m in con.execute(
        "SELECT year, meeting FROM pages WHERE meeting>0 GROUP BY year, meeting HAVING sum(length(text)>200)>=3"))
    con.close()
    NEW_MOM.mkdir(parents=True, exist_ok=True)
    got = []
    for key, v in sorted(meets.items(), key=lambda kv: kv[1]["date"]):
        yr, no = key.split("/")
        if key in have or not v["url"]:
            continue
        fn = NEW_MOM / f"محضر {no}-{yr} - {v['date']}.pdf"
        if fn.exists():
            continue
        progress(f"تحميل محضر {key} ({v['date']}) ...")
        data = http_get(v["url"], binary=True)
        if data[:4] != b"%PDF":
            continue
        fn.write_bytes(data)
        got.append(key)
    return got


# ----------------------------------------------------------------- index of minutes
def db():
    con = sqlite3.connect(DB)
    con.execute("CREATE TABLE IF NOT EXISTS files(path TEXT PRIMARY KEY, sig TEXT, pages INT, notext INT)")
    con.execute("CREATE VIRTUAL TABLE IF NOT EXISTS pages USING fts5(text, file UNINDEXED, page UNINDEXED,"
                " year UNINDEXED, meeting UNINDEXED, mdate UNINDEXED, tokenize='unicode61')")
    return con


MEET_RES = [re.compile(p) for p in (
    r"(20\d\d)\s*/\s*(\d{1,3})\s*:\s*اجتماع رقم",
    r"\(\s*(20\d\d)\s*/\s*(\d{1,3})\s*\)\s*\(\s*لاجتماع رقم",
    r"الاجتماع رقم\s*\)\s*\(\s*(20\d\d)\s*/\s*(\d{1,3})",
    r"رقم\s*\)\s*\(\s*(20\d\d)\s*/\s*(\d{1,3})")]


def pdf_files():
    fs = sorted(MOM.glob("*.pdf")) + sorted(NEW_MOM.glob("*.pdf")) if NEW_MOM.exists() else sorted(MOM.glob("*.pdf"))
    return [f for f in fs if not f.name.startswith("~")]


def rel(p):
    return str(Path(p).resolve().relative_to(ARCHIVE))


def index_minutes(progress=print, force=False):
    import pymupdf
    con = db()
    meets = load_meetings()
    known = {p: s for p, s in con.execute("SELECT path, sig FROM files")}
    current = set()
    for f in pdf_files():
        rp = rel(f)
        current.add(rp)
        st = f.stat()
        sig = f"{st.st_size}-{int(st.st_mtime)}"
        if not force and known.get(rp) == sig:
            continue
        progress(f"فهرسة {rp} ...")
        con.execute("DELETE FROM pages WHERE file=?", (rp,))
        d = pymupdf.open(f)
        fy = re.search(r"(20\d\d)", f.name)
        fyear = int(fy.group(1)) if fy else 0
        fm = re.match(r"محضر (\d+)-(\d{4})", f.name)
        cur = (int(fm.group(2)), int(fm.group(1))) if fm else (fyear, 0)
        notext, batch = 0, []
        for i in range(d.page_count):
            t = norm(d[i].get_text())
            if len(t) < 40:
                notext += 1
            if not fm:
                for rx in MEET_RES:
                    m = rx.search(t)
                    if m:
                        cur = (int(m.group(1)), int(m.group(2)))
                        break
            key = f"{cur[0]}/{cur[1]}"
            mdate = meets.get(key, {}).get("date", "")
            if not mdate and "الموافق" in t[:400]:
                m = re.search(r"م\s*(20\d\d)[/-](\d\d)[/-](\d\d)", t[:500])
                if m:
                    mdate = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
            batch.append((t, rp, i + 1, cur[0], cur[1], mdate))
            if len(batch) >= 500:
                con.executemany("INSERT INTO pages(text,file,page,year,meeting,mdate) VALUES(?,?,?,?,?,?)", batch)
                batch = []
        if batch:
            con.executemany("INSERT INTO pages(text,file,page,year,meeting,mdate) VALUES(?,?,?,?,?,?)", batch)
        # pages that inherited a meeting whose date appears only on its first page
        con.execute("INSERT OR REPLACE INTO files VALUES(?,?,?,?)", (rp, sig, d.page_count, notext))
        con.commit()
        fill_meeting_dates(con, rp)
        d.close()
    for rp in set(known) - current:
        con.execute("DELETE FROM pages WHERE file=?", (rp,))
        con.execute("DELETE FROM files WHERE path=?", (rp,))
    con.commit()
    con.close()


def fill_meeting_dates(con, rp):
    dates = {}
    for y, m, d in con.execute("SELECT year, meeting, mdate FROM pages WHERE file=? AND mdate<>''", (rp,)):
        dates.setdefault((y, m), d)
    for (y, m), d in dates.items():
        con.execute("UPDATE pages SET mdate=? WHERE file=? AND year=? AND meeting=? AND mdate=''", (d, rp, y, m))
    con.commit()


# ----------------------------------------------------------------- search: awards
_IDF = {}


def idf_table(rows):
    if not _IDF.get("n") == len(rows):
        df = collections.Counter()
        for r in rows:
            df.update(set(stems(r["subject"] + " " + r["org"])))
        n = len(rows)
        _IDF.clear()
        _IDF.update({k: math.log((n + 1) / (v + 0.5)) for k, v in df.items()})
        _IDF["n"] = n
    return _IDF


def search_awards(q, limit=40, year=None):
    rows = read_excel_rows()
    idf = idf_table(rows)
    qn = norm(q)
    qs = [s for s in dict.fromkeys(stems(q)) if not s.isdigit()]
    qnums = re.findall(r"\d+", qn)
    qflat = re.sub(r"\s", "", qn)
    wsum = sum(idf.get(s, 8.0) for s in qs) or 1.0
    out = []
    for r in rows:
        if year and not r["date"].startswith(str(year)):
            continue
        score = 0.0
        if qs:
            hay = set(stems(r["subject"] + " " + r["org"] + " " + r["winner"]))
            got = 0.0
            for s in qs:
                w = idf.get(s, 8.0)
                if s in hay:
                    got += w
                elif len(s) >= 4 and any(h.startswith(s) or s.startswith(h) for h in hay if len(h) >= 4):
                    got += 0.6 * w
            score = got / wsum
        if qnums:
            tn = num_parts(r["tno"])
            if qflat and qflat in re.sub(r"\s", "", norm(r["tno"])):
                ns = 1.0
            elif all(x in tn for x in qnums):
                ns = 0.9 if len(qnums) > 1 or len(qnums[0]) >= 5 else 0.5
            else:
                ns = 0.0
            score = max(score, ns) if not qs else 0.6 * score + 0.4 * ns + (0.3 if ns >= 0.9 else 0)
        if score >= 0.35:
            out.append((score, r["date"], r))
    out.sort(key=lambda x: (round(x[0], 2), x[1]), reverse=True)
    return [dict(r, score=round(s, 3)) for s, _d, r in out[:limit]]


def same_tender(r, rows=None):
    rows = rows or read_excel_rows()
    k = re.sub(r"\s", "", norm(r["tno"]))
    return [x for x in rows if re.sub(r"\s", "", norm(x["tno"])) == k and x["row"] != r["row"]
            and len(k) >= 4]


def similar_history(r, limit=12):
    """earlier awards with a similar subject (price history for the same kind of work)"""
    rows = read_excel_rows()
    idf = idf_table(rows)
    qs = list(dict.fromkeys(stems(r["subject"])))
    wsum = sum(idf.get(s, 8.0) for s in qs) or 1.0
    k = re.sub(r"\s", "", norm(r["tno"]))
    out = []
    for x in rows:
        if re.sub(r"\s", "", norm(x["tno"])) == k:
            continue
        hay = set(stems(x["subject"]))
        sc = sum(idf.get(s, 8.0) for s in qs if s in hay) / wsum
        if sc >= 0.55:
            out.append((sc, x["date"], x))
    out.sort(key=lambda t: (round(t[0], 1), t[1]), reverse=True)
    return [dict(x, score=round(s, 2)) for s, _d, x in out[:limit]]


# ----------------------------------------------------------------- search: minutes
TYPES = [("ترسيه", "ترسية"), ("تمديد", "تمديد"), ("امر تغييري", "أمر تغييري"), ("اوامر تغييريه", "أمر تغييري"),
         ("امر تغيير", "أمر تغييري"), ("تجديد", "تجديد"), ("الغاء", "إلغاء"), ("تعاقد مباشر", "تعاقد مباشر"),
         ("ممارسه", "ممارسة"), ("تظلم", "تظلم"), ("شكوي", "شكوى"), ("استبعاد", "استبعاد"),
         ("توصيات", "توصيات لجنة"), ("زياده", "زيادة"), ("تخفيض", "تخفيض"), ("طرح", "طرح")]


def fts_quote(w):
    return '"' + w.replace('"', "") + '"'


def num_regex(nums):
    sep = r"[\s/\-\(\)\.\\]{1,4}"
    fw = sep.join(rf"(?<!\d){n}(?!\d)" for n in nums)
    bw = sep.join(rf"(?<!\d){n}(?!\d)" for n in reversed(nums))
    return re.compile(f"{fw}|{bw}")


def search_minutes(tno="", subject="", org="", award_date="", limit=25, extra=""):
    """pages of the minutes about one tender -> list of hits (best first)"""
    if not DB.exists():
        return []
    con = db()
    nums = num_parts(tno)
    keys = [n for n in nums if not is_year(n)]
    subj_words = [w for w in dict.fromkeys(words(subject + " " + extra)) if w not in STOP and len(w) > 2]
    rows = read_excel_rows()
    idf = idf_table(rows)
    subj_words.sort(key=lambda w: -idf.get(stem(w), 8.0))
    top = subj_words[:9]
    cand = {}

    def run(q, lim):
        try:
            for rid, t, f, p, y, m, md in con.execute(
                    "SELECT rowid, text, file, page, year, meeting, mdate FROM pages WHERE pages MATCH ?"
                    " ORDER BY rank LIMIT ?", (q, lim)):
                cand[rid] = (t, f, p, y, m, md)
        except sqlite3.OperationalError:
            pass

    if nums and (keys and (len(nums) > 1 or len(keys[0]) >= 5)):
        run("NEAR(" + " ".join(fts_quote(n) for n in nums) + ", 6)", 400)
    if top:
        need = max(2, math.ceil(len(top) * 0.6))
        run(" OR ".join(fts_quote(w) for w in top), 600)
        if len(top) >= 3:
            run(" AND ".join(fts_quote(w) for w in top[:need]), 400)
    if not cand:
        return []
    qst = list(dict.fromkeys(stem(w) for w in subj_words))
    wsum = sum(idf.get(s, 8.0) for s in qst) or 1.0
    ost = [s for s in dict.fromkeys(stems(org)) if len(s) > 2]
    nrx = num_regex(nums) if len(nums) >= 2 else None
    minyear = min((int(n) for n in nums if is_year(n)), default=0)
    hits = []
    for rid, (t, f, p, y, m, md) in cand.items():
        if minyear and y and y < minyear - 1:
            continue
        pst = set(stem(w) for w in re.findall(r"[a-z0-9]+|[\u0621-\u064a]+", t))
        subj = sum(idf.get(s, 8.0) for s in qst if s in pst) / wsum if qst else 0
        pos = -1
        nsc = 0.0
        if nrx:
            mm = nrx.search(t)
            if mm and keys:
                nsc, pos = 1.0, mm.start()
        if not nsc and keys:
            if all(re.search(rf"(?<!\d){k}(?!\d)", t) for k in keys) and \
                    all(re.search(rf"(?<!\d){n}(?!\d)", t) for n in nums):
                nsc = 0.45 if len(keys[0]) < 5 else 0.9
                pos = re.search(rf"(?<!\d){keys[0]}(?!\d)", t).start()
        osc = sum(1 for s in ost if s in pst) / len(ost) if ost else 0
        date_hit = bool(award_date and md == award_date)
        if keys:
            score = 0.5 * subj + 0.38 * nsc + 0.12 * osc
        else:
            score = 0.85 * subj + 0.15 * osc
        if date_hit and subj >= 0.4:
            score += 0.25
        if score < 0.5 or subj < 0.25:
            continue
        if pos < 0 and top:
            ps = [t.find(w) for w in top[:4] if t.find(w) >= 0]
            pos = min(ps) if ps else 0
        before = t[max(0, pos - 700):pos + 150]
        typ, best = "", -1
        for k, lab in TYPES:
            i = before.rfind(k)
            if i > best:
                best, typ = i, lab
        dec = ""
        after = t[pos:pos + 2500]
        md2 = re.search(r"قرر مجلس اداره الجهاز[^.]{0,160}", after)
        if md2:
            dec = md2.group(0)
        hits.append({"id": rid, "file": f, "page": p, "year": y, "meeting": m, "mdate": md,
                     "score": round(min(score, 1.5), 3), "num": nsc, "subj": round(subj, 2),
                     "award_meeting": date_hit, "type": typ, "decision": dec[:200],
                     "snippet": t[max(0, pos - 160):pos + 340]})
    hits.sort(key=lambda h: (h["score"], h["mdate"]), reverse=True)
    # keep the strong ones, drop pages that only share generic words with the subject
    if hits:
        best = hits[0]["score"]
        hits = [h for h in hits if h["score"] >= max(0.5, best - 0.45)]
    con.close()
    return hits[:limit]


def minutes_free_search(q, limit=40):
    """direct search in the minutes with free text (tender no. and/or words)"""
    nums = re.findall(r"\d+", norm(q))
    txt = re.sub(r"[\d/\\\-]+", " ", q)
    return search_minutes(tno="/".join(nums), subject=txt, limit=limit)


# ----------------------------------------------------------------- PDF report
CSS = """
@font-face {font-family: ar; src: url(tahoma.ttf);}
@font-face {font-family: ar; src: url(tahomabd.ttf); font-weight: bold;}
* {font-family: ar; font-size: 9pt; line-height: 1.35;}
body {direction: rtl; text-align: right;}
h1 {font-size: 15pt; color: #1F4E78; margin: 0 0 4px 0; text-align: right;}
h2 {font-size: 11.5pt; color: #fff; background: #1F4E78; padding: 3px 6px; margin: 10px 0 4px 0; text-align: right;}
p {margin: 2px 0; text-align: right;}
.small {font-size: 7.5pt; color: #555;}
table {border-collapse: collapse; width: 100%;}
td, th {border: 0.6px solid #9aa7b4; padding: 2px 4px; text-align: right; vertical-align: top;}
th {background: #DDEBF7; font-weight: bold;}
.kv {font-size: 10.5pt; margin: 3px 0; border-bottom: 0.5px solid #dde3ea; padding-bottom: 2px;}
.kv b, .kv span {font-size: 10.5pt;}
.fz {background: #FFFF00;}
"""


def esc(s):
    return html.escape(str(s or ""))


def is_fawaz(s):
    return "فواز" in norm(s) or "fawaz" in norm(s)


def story_pages(html_body):
    import pymupdf
    arch = pymupdf.Archive(str(FONTS))
    story = pymupdf.Story(html=f'<body dir="rtl">{html_body}</body>', user_css=CSS, archive=arch)
    import io
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (36, 36, -36, -40)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return pymupdf.open("pdf", buf.getvalue())


def safe_name(s, n=60):
    s = re.sub(r'[\\/:*?"<>|\n\r\t]+', " ", str(s)).strip()
    return re.sub(r"\s+", " ", s)[:n].strip()


def build_report(rec, hits, query=""):
    """rec = award row (dict from Excel) or None; hits = minutes hits to include"""
    import pymupdf
    OUT.mkdir(exist_ok=True)
    today = datetime.date.today()
    parts = [f"<h1>ملف المناقصة - الجهاز المركزي للمناقصات العامة</h1>"
             f"<p class='small'>اتعمل {dmy(today)} من ملف CAPT Winning Bids from 2006 ومحاضر اجتماعات الجهاز"
             + (f" - البحث: {esc(query)}" if query else "") + "</p>"]
    if rec:
        fz = " class='fz'" if is_fawaz(rec["winner"]) else ""
        d = to_date_iso(rec["date"])
        kv = [("رقم المناقصة", esc(rec["tno"]), ""), ("الموضوع", esc(rec["subject"]), ""),
              ("الجهة", esc(rec["org"]), ""), ("الشركة الفائزة", esc(rec["winner"] or "-"), fz),
              ("القيمة الإجمالية", esc(rec["total"] or "-"), ""), ("تاريخ الترسية (الاجتماع)", dmy(d), "")]
        if rec["notes"]:
            kv.append(("ملاحظات", esc(rec["notes"]), ""))
        parts.append("<h2>بيانات الترسية</h2>" + "".join(
            f"<p class='kv'><b>{k}:</b> <span{c}>{v}</span></p>" for k, v, c in kv))
        others = same_tender(rec)
        if others:
            parts.append("<h2>صفوف تانية لنفس رقم المناقصة (فائزين/بنود/جولات تانية)</h2><table>"
                         "<tr><th>التاريخ</th><th>الفائز</th><th>القيمة</th><th>ملاحظات</th></tr>")
            for x in sorted(others, key=lambda x: x["date"], reverse=True):
                fz = " class='fz'" if is_fawaz(x["winner"]) else ""
                parts.append(f"<tr><td>{dmy(to_date_iso(x['date']))}</td><td{fz}>{esc(x['winner'])}</td>"
                             f"<td>{esc(x['total'])}</td><td>{esc(x['notes'])[:160]}</td></tr>")
            parts.append("</table>")
        hist = similar_history(rec)
        if hist:
            parts.append("<h2>ترسيات سابقة لنفس نوع الشغل (للمقارنة والتسعير)</h2><table>"
                         "<tr><th>التاريخ</th><th>رقم المناقصة</th><th>الموضوع</th><th>الجهة</th>"
                         "<th>الفائز</th><th>القيمة</th></tr>")
            for x in hist:
                fz = " class='fz'" if is_fawaz(x["winner"]) else ""
                parts.append(f"<tr><td>{dmy(to_date_iso(x['date']))}</td><td>{esc(x['tno'])}</td>"
                             f"<td>{esc(x['subject'])[:140]}</td><td>{esc(x['org'])}</td>"
                             f"<td{fz}>{esc(x['winner'])}</td><td>{esc(x['total'])}</td></tr>")
            parts.append("</table>")
    hits = sorted(hits, key=lambda h: (h.get("mdate") or "", h["page"]), reverse=True)
    if hits:
        parts.append("<h2>المرفقات من محاضر الاجتماعات (الأحدث الأول)</h2><table>"
                     "<tr><th>الاجتماع</th><th>التاريخ</th><th>نوع البند (تقريبي)</th><th>الملف / الصفحة</th>"
                     "<th>القرار</th></tr>")
        for h in hits:
            star = " (اجتماع الترسية)" if h.get("award_meeting") else ""
            parts.append(f"<tr><td>{h['year']}/{h['meeting'] or '?'}{star}</td><td>{dmy(to_date_iso(h['mdate']))}</td>"
                         f"<td>{esc(h['type'])}</td><td>{esc(Path(h['file']).name)} - ص {h['page']}</td>"
                         f"<td>{esc(h.get('decision', ''))[:120]}</td></tr>")
        parts.append("</table><p class='small'>صفحات المحاضر متضافة بعد الصفحة دي بنفس الترتيب، "
                     "والكلام المتعلق بالمناقصة متعلّم بالأصفر. نوع البند مستنتج من النص فراجعه على الصفحة.</p>")
    else:
        parts.append("<h2>المرفقات من محاضر الاجتماعات</h2><p>ملقتش صفحات محاضر مطابقة بدرجة كافية.</p>")
    out = story_pages("".join(parts))
    ncover = out.page_count
    toc = [[1, "بيانات الترسية", 1]]
    mark_words = []
    if rec:
        mark_words = [n for n in num_parts(rec["tno"]) if not is_year(n)]
    hl_stems = set()
    if rec:
        rows = read_excel_rows()
        idf = idf_table(rows)
        ss = sorted(set(stems(rec["subject"])), key=lambda s: -idf.get(s, 8.0))
        hl_stems = set(s for s in ss[:6] if len(s) >= 3)
    docs = {}
    added = set()
    for h in hits:
        for pg in (h["page"],):
            k = (h["file"], pg)
            if k in added:
                continue
            added.add(k)
            src = docs.get(h["file"])
            if src is None:
                src = docs[h["file"]] = pymupdf.open(ARCHIVE / h["file"])
            if pg - 1 >= src.page_count:
                continue
            out.insert_pdf(src, from_page=pg - 1, to_page=pg - 1)
            page = out[-1]
            highlight(page, mark_words, hl_stems, is_fawaz(rec["winner"]) if rec else False)
            label = (f"محضر {h['year']}/{h['meeting'] or '?'} - {dmy(to_date_iso(h['mdate']))}"
                     f" - {h['type'] or 'بند'} - {Path(h['file']).name} ص {pg}")
            stamp(page, label)
            toc.append([1, label, out.page_count])
    for s in docs.values():
        s.close()
    out.set_toc(toc)
    if rec:
        name = f"{safe_name(rec['tno'], 40)} - {safe_name(rec['subject'], 55)} - {rec['date'] or ''}.pdf"
    else:
        name = f"بحث المحاضر - {safe_name(query, 60)} - {today:%Y-%m-%d}.pdf"
    path = OUT / name
    try:
        out.save(path, garbage=3, deflate=True)
    except Exception:            # noqa: BLE001 - file open in a viewer
        path = OUT / (path.stem + f" ({datetime.datetime.now():%H%M%S}).pdf")
        out.save(path, garbage=3, deflate=True)
    out.close()
    return path, ncover


def to_date_iso(s):
    try:
        return datetime.date.fromisoformat(s) if s else None
    except ValueError:
        return None


def highlight(page, nums, hl_stems, fawaz):
    try:
        for w in page.get_text("words"):
            t = norm(w[4])
            hit = any(re.search(rf"(?<!\d){n}(?!\d)", t) for n in nums if n)
            if not hit and hl_stems:
                hit = stem(re.sub(r"[^\u0621-\u064a]", "", t)) in hl_stems
            if not hit and fawaz:
                hit = "فواز" in t
            if hit:
                a = page.add_highlight_annot(pymupdf_rect(w))
                a.update()
    except Exception:            # noqa: BLE001 - highlighting is best effort
        pass


def pymupdf_rect(w):
    import pymupdf
    return pymupdf.Rect(w[:4])


def stamp(page, label):
    import pymupdf
    r = page.rect
    box = pymupdf.Rect(20, 2, r.width - 20, 15)
    page.draw_rect(box, color=None, fill=(1, 0.95, 0.6), fill_opacity=0.85, overlay=True)
    page.insert_htmlbox(box, f'<div dir="rtl" style="font-size:7pt;text-align:center">{esc(label)}</div>',
                        css="* {font-family: ar;} @font-face {font-family: ar; src: url(tahoma.ttf);}",
                        archive=pymupdf.Archive(str(FONTS)))


def report_for(query, pick=0, row=None):
    rows = read_excel_rows()
    rec = None
    if row:
        rec = next((r for r in rows if r["row"] == int(row)), None)
    else:
        res = search_awards(query, limit=10)
        if res:
            rec = res[min(pick, len(res) - 1)]
    if rec:
        hits = search_minutes(rec["tno"], rec["subject"], rec["org"], rec["date"])
    else:
        hits = minutes_free_search(query)
    path, _ = build_report(rec, hits, query)
    return rec, hits, path


# ----------------------------------------------------------------- update (routine)
def run_update(progress=print):
    summary = {"new_awards": [], "fawaz": [], "minutes": [], "errors": []}
    try:
        new, last = update_excel(progress)
        summary["last_before"] = dmy(last)
        summary["new_awards"] = [{"date": dmy(r["date"]), "tno": r["tno"], "subject": r["subject"],
                                  "org": r["org"], "winner": r["winner"], "total": r["total"]} for r in new]
        summary["fawaz"] = [x for x in summary["new_awards"] if is_fawaz(x["winner"])]
        log(f"الإكسيل: {len(new)} ترسية جديدة (بعد {dmy(last)})")
    except Exception as e:          # noqa: BLE001 - report and continue with minutes
        summary["errors"].append(f"الإكسيل: {e}")
        log(f"خطأ الإكسيل: {e}")
    try:
        index_minutes(progress)     # make sure the index knows the current files first
        summary["minutes"] = download_new_minutes(progress)
        if summary["minutes"]:
            index_minutes(progress)
        log(f"المحاضر: اتحمل {len(summary['minutes'])} محضر جديد {summary['minutes']}")
    except Exception as e:          # noqa: BLE001
        summary["errors"].append(f"المحاضر: {e}")
        log(f"خطأ المحاضر: {e}")
    try:
        LAST_RUN.write_text(datetime.datetime.now().isoformat(timespec="seconds"), encoding="utf-8")
    except OSError:
        pass
    return summary


def auto_minutes():
    """minutes between automatic awards + minutes refreshes (0 = off) - the window can change it"""
    try:
        return max(0, int(cfg.S.get("awards_update_minutes", 5)))
    except (TypeError, ValueError):
        return 5


def auto_jobs(state):
    """installed program: check for a new version, and refresh awards + minutes every few minutes"""
    last_check = 0
    while True:
        if time.time() - last_check >= 1800:
            last_check = time.time()
            try:
                if updater.check_and_apply(lambda m: state["log"].append(m)):
                    os._exit(0)         # the setup takes over and restarts the programs
            except Exception:           # noqa: BLE001
                pass
        try:
            last = datetime.datetime.fromisoformat(LAST_RUN.read_text(encoding="utf-8").strip())
            age = (datetime.datetime.now() - last).total_seconds() / 60
        except (OSError, ValueError):
            age = 1e9
        every = auto_minutes()
        state["next"] = (datetime.datetime.now() + datetime.timedelta(minutes=max(0, every - age))
                         ).isoformat(timespec="minutes") if every else None
        if every and age >= every and state.get("update") != "running":
            state["update"], state["log"] = "running", ["تحديث تلقائي للترسيات والمحاضر"]
            try:
                s = run_update(lambda m: state["log"].append(m))
                state["log"].append(f"خلص: {len(s['new_awards'])} ترسية جديدة، {len(s['minutes'])} محضر جديد")
                state["update"] = "done"
            except Exception as e:      # noqa: BLE001
                state["log"].append(f"خطأ: {e}")
                state["update"] = "error"
        time.sleep(20)


def index_stats():
    if not DB.exists():
        return {}
    con = db()
    n = con.execute("SELECT count(*) FROM pages").fetchone()[0]
    nt = con.execute("SELECT sum(notext), count(*) FROM files").fetchone()
    last = con.execute("SELECT year, meeting, mdate FROM pages WHERE mdate<>'' ORDER BY mdate DESC LIMIT 1").fetchone()
    con.close()
    return {"pages": n, "notext": nt[0] or 0, "files": nt[1], "last_meeting": list(last) if last else None}


# ----------------------------------------------------------------- web UI
def serve(open_browser=True):
    from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
    import urllib.parse
    import webbrowser
    state = {"update": None, "log": []}

    class H(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def send(self, obj, code=200, ctype="application/json; charset=utf-8"):
            data = obj if isinstance(obj, bytes) else json.dumps(obj, ensure_ascii=False, default=str).encode()
            self.send_response(code)
            self.send_header("Content-Type", ctype)
            self.send_header("Content-Length", str(len(data)))
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            self.wfile.write(data)

        def do_GET(self):
            u = urllib.parse.urlparse(self.path)
            q = {k: v[0] for k, v in urllib.parse.parse_qs(u.query).items()}
            try:
                if u.path == "/":
                    return self.send((cfg.res_dir() / "ui" / "واجهة البرنامج.html").read_bytes(), ctype="text/html; charset=utf-8")
                if u.path == "/api/status":
                    rows = read_excel_rows()
                    last = max((r["date"] for r in rows if r["date"]), default="")
                    return self.send({"rows": len(rows), "last_award": last, "index": index_stats(),
                                      "version": cfg.build_info().get("version"),
                                      "update": state["update"], "log": state["log"][-15:],
                                      "auto_minutes": auto_minutes(), "next": state.get("next"),
                                      "last_update": LAST_RUN.read_text(encoding="utf-8").strip()
                                      if LAST_RUN.exists() else ""})
                if u.path == "/api/search":
                    return self.send(search_awards(q.get("q", ""), limit=int(q.get("limit", 40)),
                                                   year=q.get("year") or None))
                if u.path == "/api/detail":
                    rows = read_excel_rows()
                    rec = next(r for r in rows if r["row"] == int(q["row"]))
                    return self.send({"rec": rec, "same": same_tender(rec), "history": similar_history(rec),
                                      "hits": search_minutes(rec["tno"], rec["subject"], rec["org"], rec["date"])})
                if u.path == "/api/minutes":
                    return self.send(minutes_free_search(q.get("q", "")))
                if u.path == "/api/page.png":
                    import pymupdf
                    d = pymupdf.open(ARCHIVE / q["file"])
                    png = d[int(q["page"]) - 1].get_pixmap(dpi=int(q.get("dpi", 110))).tobytes("png")
                    d.close()
                    return self.send(png, ctype="image/png")
                if u.path == "/report":
                    p = OUT / Path(q["name"]).name
                    return self.send(p.read_bytes(), ctype="application/pdf")
                self.send({"error": "not found"}, 404)
            except Exception as e:          # noqa: BLE001
                self.send({"error": str(e)}, 500)

        def do_POST(self):
            u = urllib.parse.urlparse(self.path)
            body = json.loads(self.rfile.read(int(self.headers.get("Content-Length", 0)) or 0) or b"{}")
            try:
                if u.path == "/api/report":
                    rows = read_excel_rows()
                    rec = next((r for r in rows if r["row"] == int(body["row"])), None) if body.get("row") else None
                    path, _ = build_report(rec, body.get("hits", []), body.get("query", ""))
                    if body.get("open", True):
                        os.startfile(path)
                    return self.send({"path": str(path), "name": path.name})
                if u.path == "/api/open-folder":
                    OUT.mkdir(exist_ok=True)
                    os.startfile(OUT)
                    return self.send({"ok": True})
                if u.path == "/api/open-excel":
                    os.startfile(XLSX)
                    return self.send({"ok": True})
                if u.path == "/api/update":
                    if state["update"] == "running":
                        return self.send({"status": "running"})
                    state["update"], state["log"] = "running", []

                    def job():
                        try:
                            s = run_update(lambda m: state["log"].append(m))
                            state["log"].append(f"خلص: {len(s['new_awards'])} ترسية جديدة، "
                                                f"{len(s['minutes'])} محضر جديد" +
                                                (f" - أخطاء: {s['errors']}" if s["errors"] else ""))
                            state["update"] = "done"
                        except Exception as e:      # noqa: BLE001
                            state["log"].append(f"خطأ: {e}")
                            state["update"] = "error"
                    threading.Thread(target=job, daemon=True).start()
                    return self.send({"status": "running"})
                if u.path == "/api/settings":
                    if "awards_update_minutes" in body:
                        cfg.save_setting("awards_update_minutes", max(0, int(body["awards_update_minutes"])))
                    return self.send({"auto_minutes": auto_minutes()})
                self.send({"error": "not found"}, 404)
            except Exception as e:          # noqa: BLE001
                self.send({"error": str(e)}, 500)

    class Server(ThreadingHTTPServer):
        allow_reuse_address = False     # on Windows reuse lets a 2nd copy bind the same port

    url = f"http://127.0.0.1:{PORT}/"
    try:
        srv = Server(("127.0.0.1", PORT), H)
    except OSError:                 # already running -> just open it
        if open_browser:
            webbrowser.open(url)
        return
    threading.Thread(target=auto_jobs, args=(state,), daemon=True).start()
    if open_browser:
        threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    srv.serve_forever()


# ----------------------------------------------------------------- CLI
def _print_award(i, r):
    print(f"{i}) [{r['date']}] {r['tno']} | {r['subject'][:110]} | {r['org']} | فاز: {r['winner'] or '-'}"
          f" | {r['total'] or '-'} | صف {r['row']} | score {r.get('score', '')}")


def main(argv):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if not argv:
        print(__doc__)
        return
    cmd, rest = argv[0], argv[1:]
    opt = {}
    args = []
    it = iter(rest)
    for a in it:
        if a.startswith("--"):
            opt[a[2:]] = next(it, "1")
        else:
            args.append(a)
    q = " ".join(args)
    if cmd == "update":
        s = run_update()
        print(json.dumps(s, ensure_ascii=False, indent=1))
    elif cmd == "index":
        index_minutes(force="force" in opt)
        print(json.dumps(index_stats(), ensure_ascii=False))
    elif cmd == "search":
        for i, r in enumerate(search_awards(q, limit=int(opt.get("limit", 15)))):
            _print_award(i, r)
    elif cmd == "minutes":
        for h in minutes_free_search(q):
            print(f"{h['year']}/{h['meeting']} {h['mdate']} {h['file']} ص{h['page']} score {h['score']} {h['type']}")
            print("    ", h["snippet"][:300])
    elif cmd == "report":
        rec, hits, path = report_for(q, pick=int(opt.get("pick", 0)), row=opt.get("row"))
        if rec:
            _print_award("*", rec)
        print(f"صفحات محاضر: {len(hits)}")
        for h in sorted(hits, key=lambda h: h["mdate"], reverse=True):
            print(f"   {h['year']}/{h['meeting']} {h['mdate']} {h['type']} | {h['file']} ص{h['page']} | score {h['score']}")
        print("PDF:", path)
    elif cmd == "serve":
        serve(open_browser="nobrowser" not in opt)
    else:
        print(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
